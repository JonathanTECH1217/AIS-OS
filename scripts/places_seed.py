"""Harvest a national seed list of home integrators, AV companies, and home theater builders from the
Google Places API (New) Text Search, and write it in the Integrator_List workbook layout.

Grid: scripts/places_grid.py (metros get all QUERIES, enclaves the first four). Each query pages up to
3 times (Google caps text search at 60 results) and stops early when a page adds fewer than 5 new places
or under 20% new. Every page is cached in projects/outreach/places-cache-<DATE>.jsonl so a rerun never
re-buys a finished query.

Billing: the field mask includes websiteUri, phone, and rating, so every page is a Text Search Enterprise
request, $35.00 per 1,000 (first 1,000 per month free at the project level). Each pageToken call bills
again. The CostMeter refuses to start a page that would cross --max-cost.

Dedupe sources (read only, never written): projects/outreach/Integrator_List_2026-09-05.xlsx (Extras sheet
has Place IDs) and, when the sibling repo is on this machine, monarc-os/Data/Companies.xlsx (Integrator rows).
Those rows are merged into the seed with "In Repository" = yes so nothing is bought twice.

Usage:
  python scripts/places_seed.py --check-keys
  python scripts/places_seed.py --dry-run [--limit-locations N]
  python scripts/places_seed.py [--limit-locations N] [--max-cost USD] [--queries "a;b"] [--date YYYY-MM-DD]
  python scripts/places_seed.py --build-only          # rebuild the workbook from the cache

Output: projects/outreach/places-seed-<DATE>.xlsx (sheets Integrators, Extras, Method) and a JSON summary.
Key: GOOGLE_MAPS_API_KEY, resolved by scripts/outreach_common.load_env (AIOS .env, then
%USERPROFILE%/.monarc/secrets.env, then the environment). See references/google-places-api.md.
bike-method-phase: 1
"""
import argparse
import json
import re
import sys
import time
from datetime import date
from pathlib import Path

import openpyxl
import requests
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from outreach_common import (  # noqa: E402
    DMV_STATES, OUTREACH, ROOT, BudgetExceeded, CostMeter, budget_cap, get_secret, local_time_noon_et,
    norm_domain, parse_addr, report_key_sources, reviews_string,
)
from places_grid import ENCLAVES, LOCATIONS, QUERIES  # noqa: E402

SEARCH_URL = "https://places.googleapis.com/v1/places:searchText"
FIELD_MASK = ",".join([
    "places.id", "places.displayName", "places.formattedAddress", "places.addressComponents",
    "places.nationalPhoneNumber", "places.websiteUri", "places.rating", "places.userRatingCount",
    "places.businessStatus", "places.primaryType", "places.types", "places.googleMapsUri", "nextPageToken",
])
PRICE_PER_REQUEST = 35.00 / 1000.0
MAX_PAGES = 3
PAGE_SLEEP = 2.0
QUERY_SLEEP = 0.25
EARLY_STOP_MIN_NEW = 5
EARLY_STOP_RATIO = 0.20

EXISTING_XLSX = OUTREACH / "Integrator_List_2026-09-05.xlsx"
MONARC_OS_XLSX = ROOT.parent / "monarc-os" / "Data" / "Companies.xlsx"

INTEGRATOR_COLS = ["Name", "Address", "State", "County", "Phone", "Local Time @ 12pm ET", "Reviews", "Website",
                   "Status", "Outcome"]
EXTRA_COLS = ["Place ID", "Google Maps URL", "Primary type", "Business status", "Domain", "Source list", "DMV",
              "City", "Zip", "Rating", "Review count", "Query"]
EXTRAS_SHEET_COLS = ["Name", "Website", "Rating", "Reviews", "Place ID", "Query", "In Repository"]


# ---------------------------------------------------------------- Places call

def search_page(api_key, query, page_token=None):
    body = {"textQuery": query, "regionCode": "US", "languageCode": "en", "pageSize": 20}
    if page_token:
        body["pageToken"] = page_token
    headers = {"Content-Type": "application/json", "X-Goog-Api-Key": api_key, "X-Goog-FieldMask": FIELD_MASK}
    last = None
    for attempt in range(5):
        try:
            r = requests.post(SEARCH_URL, json=body, headers=headers, timeout=25)
        except requests.RequestException as e:
            last = type(e).__name__
            time.sleep(2 * (attempt + 1))
            continue
        if r.status_code == 200:
            data = r.json()
            return data.get("places", []), data.get("nextPageToken"), None
        if r.status_code in (403, 429, 500, 502, 503, 504):
            # 403 is retried too (2026-09-24): after a key or API change in the console, Google's servers answer
            # "The caller does not have permission" on and off for a few minutes while the change settles.
            last = f"HTTP {r.status_code}"
            time.sleep((10 if r.status_code == 403 else 3) * (attempt + 1))
            continue
        try:
            msg = r.json().get("error", {}).get("message", r.text[:200])
        except ValueError:
            msg = r.text[:200]
        return [], None, f"HTTP {r.status_code}: {msg}"
    return [], None, last or "request failed"


def component(place, kind):
    for c in place.get("addressComponents", []) or []:
        if kind in (c.get("types") or []):
            return c.get("shortText") or c.get("longText") or ""
    return ""


def flatten(place, query):
    addr = place.get("formattedAddress", "") or ""
    city, state, zipc = parse_addr(addr)
    state = component(place, "administrative_area_level_1") or state
    city = component(place, "locality") or city
    zipc = component(place, "postal_code") or zipc
    county = component(place, "administrative_area_level_2")
    name = (place.get("displayName") or {}).get("text", "") or ""
    return {
        "Name": name, "Address": addr, "State": state, "County": county,
        "Phone": place.get("nationalPhoneNumber", "") or "",
        "Website": place.get("websiteUri", "") or "",
        "Rating": place.get("rating"), "Review count": place.get("userRatingCount") or 0,
        "Place ID": place.get("id", ""), "Google Maps URL": place.get("googleMapsUri", "") or "",
        "Primary type": place.get("primaryType", "") or "", "Business status": place.get("businessStatus", "") or "",
        "City": city, "Zip": zipc, "Query": query,
    }


# ---------------------------------------------------------------- cache

def cache_path(run_date):
    return OUTREACH / f"places-cache-{run_date}.jsonl"


def load_cache(path):
    """Returns (records_by_place_id, finished_queries, pages_bought)."""
    recs, finished, pages = {}, set(), 0
    if not path.exists():
        return recs, finished, pages
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        pages += 1
        for p in row.get("places", []):
            recs.setdefault(p["Place ID"], p)
        if row.get("complete"):
            finished.add(row["query"])
    return recs, finished, pages


# ---------------------------------------------------------------- existing lists

def load_existing_integrator_list(path=EXISTING_XLSX):
    out = {}
    if not path.exists():
        return out
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb["Integrators"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h) for h in next(rows)]
    ix = {h: i for i, h in enumerate(hdr)}
    main = []
    for r in rows:
        if r[ix["Name"]]:
            main.append(r)
    extras, ix2 = {}, {}
    if "Extras" in wb.sheetnames:
        ws2 = wb["Extras"]
        rows2 = ws2.iter_rows(values_only=True)
        hdr2 = [str(h) for h in next(rows2)]
        ix2 = {h: i for i, h in enumerate(hdr2)}
        for r in rows2:
            if r[ix2["Place ID"]]:
                extras[(str(r[ix2["Name"]]), str(r[ix2["Website"]] or ""))] = r
    for r in main:
        name, site = str(r[ix["Name"]]), str(r[ix["Website"]] or "")
        ex = extras.get((name, site))
        pid = str(ex[ix2["Place ID"]]) if ex else f"legacy:{name}|{site}"
        rating = ex[ix2["Rating"]] if ex else None
        count = ex[ix2["Reviews"]] if ex else None
        if rating is None:
            m = re.match(r"\s*([\d.]+)", str(r[ix["Reviews"]] or ""))
            rating = float(m.group(1)) if m else None
        if count is None:
            m = re.search(r"\((\d+)\)", str(r[ix["Reviews"]] or ""))
            count = int(m.group(1)) if m else 0
        city, _, zipc = parse_addr(str(r[ix["Address"]] or ""))
        out[pid] = {
            "Name": name, "Address": r[ix["Address"]] or "", "State": r[ix["State"]] or "",
            "County": r[ix["County"]] or "", "Phone": r[ix["Phone"]] or "", "Website": site,
            "Rating": rating, "Review count": int(count or 0), "Place ID": pid, "Google Maps URL": "",
            "Primary type": "", "Business status": "", "City": city, "Zip": zipc,
            "Query": str(ex[ix2["Query"]]) if ex else "", "Source list": "2026-09-05",
            "Status": r[ix["Status"]] or "", "Outcome": r[ix["Outcome"]] or "",
        }
    return out


def load_monarc_os_integrators(path=MONARC_OS_XLSX):
    out = {}
    if not path.exists():
        return out
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb["Companies"]
    except Exception as e:  # noqa: BLE001
        print(f"  monarc-os Companies.xlsx not readable ({type(e).__name__}); skipping", flush=True)
        return out
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h) for h in next(rows)]
    ix = {h: i for i, h in enumerate(hdr)}
    for r in rows:
        if (r[ix["Vertical"]] or "") != "Integrator" or not r[ix["Place ID"]]:
            continue
        addr = r[ix["Address"]] or ""
        city, st, zipc = parse_addr(str(addr))
        out[str(r[ix["Place ID"]])] = {
            "Name": r[ix["Company"]] or "", "Address": addr, "State": r[ix["State"]] or st,
            "County": r[ix["County"]] or "", "Phone": r[ix["Phone"]] or "", "Website": r[ix["Website"]] or "",
            "Rating": r[ix["Rating"]], "Review count": int(r[ix["Review Count"]] or 0),
            "Place ID": str(r[ix["Place ID"]]), "Google Maps URL": r[ix["Google Maps URL"]] or "",
            "Primary type": r[ix["Primary Type"]] or "", "Business status": r[ix["Business Status"]] or "",
            "City": city, "Zip": zipc, "Query": "", "Source list": "monarc-os", "Status": "", "Outcome": "",
        }
    return out


# ---------------------------------------------------------------- plan and sweep

def plan_queries(locations, queries):
    enclaves = {e.lower() for e in ENCLAVES}
    plan = []
    for loc in locations:
        qs = queries[:4] if loc.lower() in enclaves and len(queries) > 4 else queries
        for q in qs:
            plan.append(q.format(loc=loc))
    return plan


def sweep(api_key, plan, run_date, meter, known_ids):
    path = cache_path(run_date)
    recs, finished, pages_before = load_cache(path)
    todo = [q for q in plan if q not in finished]
    print(f"{len(plan)} queries planned, {len(finished)} finished in cache, {len(todo)} to run; "
          f"{len(known_ids)} place IDs already known", flush=True)
    seen = set(known_ids) | set(recs)
    bought = 0
    errors = []
    with path.open("a", encoding="utf-8") as fh:
        for i, q in enumerate(todo, 1):
            token = None
            for page in range(1, MAX_PAGES + 1):
                meter.assert_affordable(PRICE_PER_REQUEST, f"Places page {page} of '{q}'")
                places, next_token, err = search_page(api_key, q, token)
                meter.add("places_requests", PRICE_PER_REQUEST)
                bought += 1
                if err:
                    errors.append({"query": q, "page": page, "error": err})
                    print(f"  ! {q} p{page}: {err}", flush=True)
                    if err.startswith("HTTP 4") or err == "HTTP 403":
                        raise SystemExit(f"Places API refused the request: {err}. Fix the key or quota and rerun.")
                    break
                flat = [flatten(p, q) for p in places if p.get("id")]
                new = [p for p in flat if p["Place ID"] not in seen]
                for p in flat:
                    seen.add(p["Place ID"])
                    recs.setdefault(p["Place ID"], p)
                complete = (not next_token) or page == MAX_PAGES or len(new) < EARLY_STOP_MIN_NEW \
                    or (len(flat) and len(new) / len(flat) < EARLY_STOP_RATIO)
                fh.write(json.dumps({"query": q, "page": page, "n": len(flat), "new": len(new),
                                     "complete": bool(complete), "places": flat,
                                     "ts": time.strftime("%Y-%m-%d %H:%M")}) + "\n")
                fh.flush()
                if complete:
                    break
                token = next_token
                time.sleep(PAGE_SLEEP)
            if i % 25 == 0 or i == len(todo):
                print(f"  {i}/{len(todo)} queries, {bought} pages bought this run, {len(recs)} places cached, "
                      f"spend ${meter.total:.2f}", flush=True)
            time.sleep(QUERY_SLEEP)
    return recs, bought, errors


# ---------------------------------------------------------------- workbook

def build(recs, existing, run_date, summary_extra):
    merged = {}
    for pid, r in existing.items():
        merged[pid] = {**r, "In Repository": "yes"}
    new_count = 0
    for pid, r in recs.items():
        if pid in merged:
            continue
        merged[pid] = {**r, "Source list": f"places-{run_date}", "Status": "", "Outcome": "", "In Repository": "no"}
        new_count += 1

    rows = []
    closed = 0
    for r in merged.values():
        if (r.get("Business status") or "") == "CLOSED_PERMANENTLY":
            closed += 1
            continue
        r["Domain"] = norm_domain(r.get("Website"))
        r["DMV"] = "yes" if (r.get("State") or "").upper() in DMV_STATES else "no"
        r["Local Time @ 12pm ET"] = local_time_noon_et(r.get("State"))
        r["Reviews"] = reviews_string(r.get("Rating"), r.get("Review count"))
        rows.append(r)
    rows.sort(key=lambda r: ((r.get("State") or ""), (r.get("City") or ""), (r.get("Name") or "")))

    out = OUTREACH / f"places-seed-{run_date}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Integrators"
    cols = INTEGRATOR_COLS + EXTRA_COLS
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") if r.get(c) is not None else "" for c in cols])
    ws2 = wb.create_sheet("Extras")
    ws2.append(EXTRAS_SHEET_COLS)
    for r in rows:
        ws2.append([r.get("Name", ""), r.get("Website", ""), r.get("Rating") if r.get("Rating") is not None else "",
                    r.get("Review count", 0), r.get("Place ID", ""), r.get("Query", "") or "(already in repository)",
                    r.get("In Repository", "no")])
    ws3 = wb.create_sheet("Method")
    for line in [
        f"Built {run_date} by scripts/places_seed.py from the Places API (New) Text Search.",
        f"Grid: {summary_extra.get('locations', '?')} locations x up to {len(QUERIES)} queries "
        f"(enclaves get the first 4), max {MAX_PAGES} pages per query, early stop when a page adds "
        f"fewer than {EARLY_STOP_MIN_NEW} new places or under {int(EARLY_STOP_RATIO * 100)}% new.",
        "One row per Google Place ID. Duplicate listings for one website are collapsed later by qualify_list.py "
        "and qualify_sites.py (highest-reviewed listing represents the company).",
        "Merged read-only from Integrator_List_2026-09-05.xlsx (Source list 2026-09-05) and, when present, "
        "monarc-os Data/Companies.xlsx Integrator rows (Source list monarc-os). In Repository = yes for those.",
        f"Permanently closed listings dropped: {closed}.",
        "DMV = yes marks DC, MD, VA. The LinkedIn ad seed keeps them (decision 2026-09-13); "
        "qualify_list.py still excludes them from the dial list.",
        "Reviews = 'rating (count)'. Local Time @ 12pm ET derived from state.",
        f"Places pages bought (all runs this date): {summary_extra.get('pages_total', '?')}; "
        f"estimated spend ${summary_extra.get('spend_usd', 0):.2f} before the monthly free tier.",
        "Cache: places-cache-<DATE>.jsonl. Rebuild without spending: python scripts/places_seed.py --build-only",
    ]:
        ws3.append([line])
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
    wb.save(out)

    by_state, by_source = {}, {}
    for r in rows:
        by_state[r.get("State") or "?"] = by_state.get(r.get("State") or "?", 0) + 1
        by_source[r.get("Source list")] = by_source.get(r.get("Source list"), 0) + 1
    domains = {r["Domain"] for r in rows if r["Domain"]}
    return out, {
        "rows": len(rows), "new_from_places": new_count, "closed_dropped": closed,
        "unique_domains": len(domains), "no_website": sum(1 for r in rows if not r["Domain"]),
        "dmv_rows": sum(1 for r in rows if r["DMV"] == "yes"),
        "by_source": by_source, "top_states": dict(sorted(by_state.items(), key=lambda kv: -kv[1])[:12]),
    }


# ---------------------------------------------------------------- main

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--check-keys", action="store_true", help="show which source answers for each key, then exit")
    ap.add_argument("--dry-run", action="store_true", help="plan and price the sweep, make no calls")
    ap.add_argument("--build-only", action="store_true", help="rebuild the workbook from the cache")
    ap.add_argument("--limit-locations", type=int, default=0, help="only the first N locations (smoke test)")
    ap.add_argument("--max-cost", type=float, default=None,
                    help="USD ceiling for this script (default: 45%% of SEED_BUDGET_USD, itself default 200)")
    ap.add_argument("--queries", default=None, help="semicolon-separated templates with {loc}, overrides the grid")
    ap.add_argument("--date", default=date.today().isoformat(), help="run date for cache and output names")
    a = ap.parse_args(argv)

    if a.check_keys:
        ok = report_key_sources(["GOOGLE_MAPS_API_KEY", "ANTHROPIC_API_KEY"])
        return 0 if ok else 1

    queries = [q.strip() for q in a.queries.split(";") if q.strip()] if a.queries else QUERIES
    locations = LOCATIONS[:a.limit_locations] if a.limit_locations else LOCATIONS
    plan = plan_queries(locations, queries)
    cap = a.max_cost if a.max_cost is not None else round(budget_cap() * 0.45, 2)
    _, finished, pages_before = load_cache(cache_path(a.date))
    todo = [q for q in plan if q not in finished]

    if a.dry_run:
        low, high = len(todo), len(todo) * MAX_PAGES
        print(json.dumps({
            "locations": len(locations), "queries_per_metro": len(queries), "text_searches": len(plan),
            "already_finished_in_cache": len(plan) - len(todo), "requests_low": low, "requests_high": high,
            "cost_low_usd": round(low * PRICE_PER_REQUEST, 2), "cost_high_usd": round(high * PRICE_PER_REQUEST, 2),
            "note": "first 1,000 Enterprise requests per month are free at the Google project level",
            "cap_for_this_script_usd": cap, "cache": str(cache_path(a.date)),
        }, indent=2))
        return 0

    print("Loading existing lists for dedupe...", flush=True)
    existing = load_existing_integrator_list()
    monarc = load_monarc_os_integrators()
    for pid, r in monarc.items():
        existing.setdefault(pid, r)
    print(f"  {len(existing)} known rows ({sum(1 for r in existing.values() if r['Source list'] == 'monarc-os')} "
          f"from monarc-os)", flush=True)

    meter = CostMeter(cap, {"places_requests": round(pages_before * PRICE_PER_REQUEST, 4)})
    bought, errors = 0, []
    if a.build_only:
        recs, _, _ = load_cache(cache_path(a.date))
    else:
        api_key = get_secret("GOOGLE_MAPS_API_KEY", required=True,
                             hint="Copy it from the Monarc OS main machine into %USERPROFILE%\\.monarc\\secrets.env.")
        try:
            recs, bought, errors = sweep(api_key, plan, a.date, meter, set(existing))
        except BudgetExceeded as e:
            print(f"STOPPED: {e}", flush=True)
            recs, _, _ = load_cache(cache_path(a.date))
    _, _, pages_total = load_cache(cache_path(a.date))
    out, summary = build(recs, existing, a.date, {"locations": len(locations), "pages_total": pages_total,
                                                   "spend_usd": pages_total * PRICE_PER_REQUEST})
    print(json.dumps({**summary, "pages_bought_this_run": bought, "pages_total_this_date": pages_total,
                      "errors": errors[:10], "cost": meter.summary(), "xlsx": str(out)}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
