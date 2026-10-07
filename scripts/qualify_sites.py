"""Qualify every unique website in a seed workbook for the LinkedIn company-list audience.

Three passes, all resumable through projects/outreach/site-cache-<DATE>.jsonl:
  1. Fetch: homepage plus up to two same-domain pages (services, residential, about, solutions). 24 threads.
  2. Platform: deterministic fingerprints for WordPress, Wix, Squarespace, GoDaddy, Shopify, Weebly, Duda,
     Webflow, HubSpot; One Firefly agency detection. No model involved.
  3. Opus 5 read (claude-opus-5, structured JSON output, effort low) through the Message Batches API at half
     price: is it an integrator, residential or commercial, which offerings, which ICP brands, site quality.
Then scores each company (Tier A / B / C / exclude) and writes:
  projects/outreach/qualified-seed-<DATE>.xlsx   (Qualified, Excluded, Unverified, All scored, Method)
  projects/outreach/linkedin-company-list-<DATE>.csv          (Tier A + B, DMV rows kept, decision 2026-09-13)
  projects/outreach/linkedin-company-list-<DATE>-no-dmv.csv   (same minus DC, MD, VA, for the dial-aligned view)

Usage:
  python scripts/qualify_sites.py [--input seed.xlsx] [--limit N] [--dry-run] [--no-batch] [--build-only]
                                  [--max-cost USD] [--date YYYY-MM-DD] [--no-wait]
  --dry-run   fetch and fingerprint, skip the model, print the Opus cost estimate
  --no-batch  synchronous messages.create (full price, instant) for a small --limit smoke test
  --no-wait   submit the batch and exit; rerun later to collect results
Keys: ANTHROPIC_API_KEY via scripts/outreach_common.load_env. Budget: SEED_BUDGET_USD (default 200) minus what
places_seed.py spent on the same date. See references/anthropic-api.md.
bike-method-phase: 1
"""
import argparse
import json
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from pathlib import Path
from urllib.parse import urljoin, urlparse

import openpyxl
from bs4 import BeautifulSoup
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from outreach_common import (  # noqa: E402
    DMV_STATES, INTEGRATOR_RE, ONEFIREFLY_RE, OUTREACH, ROOT, BudgetExceeded, CostMeter, budget_cap, clean_name,
    fetch, get_secret, newest, norm_domain, parse_addr, write_linkedin_csv,
)

MODEL = "claude-opus-5"
BATCH_IN_PER_M, BATCH_OUT_PER_M = 2.50, 12.50
SYNC_IN_PER_M, SYNC_OUT_PER_M = 5.00, 25.00
CACHE_READ_PER_M = 0.50
EST_OUTPUT_TOKENS = 1000          # JSON answer plus adaptive thinking at effort low
MAX_TOKENS = 3000
PLACES_PRICE_PER_REQUEST = 35.00 / 1000.0
HOME_TEXT_CAP = 5000
SUB_TEXT_CAP = 2000
MAX_SUBPAGES = 2
THIN_TEXT = 300                   # under this with no subpages: JS-rendered or blocked, not worth a model call
BATCH_CHUNK = 1000
POLL_SECONDS = 60
THREADS = 24

SUBPAGE_RE = re.compile(r"service|residential|about|what-we-do|solutions|smart-home|home-automation|theater|"
                        r"lighting|shades|audio", re.I)
WEAK_PLATFORMS = {"wordpress", "wix", "squarespace", "godaddy", "shopify", "weebly", "duda"}

# (platform, html regexes, header regexes). First match wins; order matters (Wix before WordPress, etc.).
PLATFORM_RULES = [
    ("wix", [r"static\.wixstatic\.com", r"static\.parastorage\.com", r"wix\.com/website", r'generator" content="Wix'],
     [r"^x-wix-"]),
    ("squarespace", [r"static1\.squarespace\.com", r"squarespace\.com", r'generator" content="Squarespace'],
     [r"squarespace"]),
    ("godaddy", [r"img1\.wsimg\.com", r"secureserver\.net", r'generator" content="(?:Starfield|GoDaddy)',
                 r"godaddy\.com/websites", r"wsimg\.com"], [r"^x-godaddy", r"secureserver"]),
    ("shopify", [r"cdn\.shopify\.com", r"Shopify\.theme", r"myshopify\.com"], [r"^x-shopify"]),
    ("weebly", [r"weebly\.com", r"editmysite\.com"], []),
    ("duda", [r"cdn\.multiscreensite\.com", r"dudamobile", r"irp\.cdn-website\.com", r"lirp\.cdn-website\.com"],
     [r"duda"]),
    ("webflow", [r"assets\.website-files\.com", r'generator" content="Webflow', r"webflow\.com"], []),
    ("hubspot", [r"hs-scripts\.com", r"hubspotusercontent", r"hs-sites\.com"], [r"hubspot"]),
    ("wordpress", [r"/wp-content/", r"/wp-includes/", r'generator" content="WordPress', r"wp-json"], [r"wordpress"]),
]
PLATFORM_LABELS = {"wordpress": "WordPress", "wix": "Wix", "squarespace": "Squarespace", "godaddy": "GoDaddy",
                   "shopify": "Shopify", "weebly": "Weebly", "duda": "Duda", "webflow": "Webflow",
                   "hubspot": "HubSpot", "custom": "custom/unknown"}

SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["is_integrator", "market_focus", "offerings", "brands", "high_end_signals", "site_quality",
                 "confidence", "reason"],
    "properties": {
        "is_integrator": {"type": "boolean",
                          "description": "True if the company designs or installs residential technology: "
                                         "control/automation, AV, home theater, lighting control, shades, "
                                         "networking, or low-voltage. False for IT shops, alarm-only companies, "
                                         "retailers, electricians without AV, gate or garage companies."},
        "market_focus": {"type": "string",
                         "enum": ["residential", "mixed_residential_lead", "mixed_commercial_lead", "commercial",
                                  "unclear"]},
        "offerings": {
            "type": "object", "additionalProperties": False,
            "required": ["lighting_control", "automation_control", "audio_video", "home_theater",
                         "motorized_shades", "networking", "security_cameras", "outdoor_av"],
            "properties": {k: {"type": "boolean"} for k in
                           ["lighting_control", "automation_control", "audio_video", "home_theater",
                            "motorized_shades", "networking", "security_cameras", "outdoor_av"]},
        },
        "brands": {"type": "array", "items": {"type": "string"},
                   "description": "ICP brands named on the pages, spelled as in the ICP list. Empty if none."},
        "high_end_signals": {"type": "array", "items": {"type": "string"},
                             "description": "Any of: Crestron, Savant, Josh.ai, Lutron HomeWorks, Kaleidescape, "
                                            "Steinway Lyngdorf, HTSA, ProSource, or a short phrase of estate or "
                                            "luxury language quoted from the page."},
        "site_quality": {
            "type": "object", "additionalProperties": False,
            "required": ["has_service_pages", "has_service_area_pages", "has_gallery", "has_testimonials",
                         "single_page"],
            "properties": {k: {"type": "boolean"} for k in
                           ["has_service_pages", "has_service_area_pages", "has_gallery", "has_testimonials",
                            "single_page"]},
        },
        "confidence": {"type": "number", "description": "0 to 1. Below 0.6 when the pages are thin or ambiguous."},
        "reason": {"type": "string", "description": "One sentence, under 200 characters, plain facts from the page."},
    },
}


def build_system_prompt():
    icp = (ROOT / "context" / "icp-brands.md").read_text(encoding="utf-8")
    brands = icp.split("\n---")[0]
    return (
        "You qualify prospects for Monarc Build, a marketing agency that sells managed Google Ads, SEO, and web "
        "management to home integrators (low voltage, AV, smart home) doing residential work in high profile homes "
        "($2M to $25M). You read extracted website text and return one JSON object matching the schema. "
        "Judge only from the text given. Do not guess brands that are not on the page. A company that serves both "
        "homes and businesses is mixed; decide which leads from the headline, navigation, and word share. "
        "Commercial-only signals: conference rooms, digital signage, houses of worship, education, hospitality, "
        "pro AV, enterprise, video conferencing, with no homeowner language. Residential signals: homeowner, home "
        "theater, smart home, whole-home audio, media room, custom homes, builders and architects and interior "
        "designers as partners, motorized shades, luxury or estate language.\n\n"
        "Reference: the ICP brand map (verbatim from Jonathan Beach).\n\n" + brands.strip() + "\n\n"
        "High-end tier rule: a Crestron, Savant, Josh.ai, Lutron HomeWorks, Kaleidescape, or Steinway Lyngdorf dealer, "
        "or an HTSA or ProSource member, is more likely doing the $2M to $25M homes Monarc targets."
    )


# ---------------------------------------------------------------- seed loading

def load_seed(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h) for h in next(rows)]
    ix = {h: i for i, h in enumerate(hdr)}

    def col(r, name, default=""):
        i = ix.get(name)
        return (r[i] if i is not None and i < len(r) and r[i] is not None else default)

    recs = []
    for r in rows:
        if not col(r, "Name"):
            continue
        site = str(col(r, "Website")).strip()
        rating, count = col(r, "Rating", None), col(r, "Review count", None)
        if count in ("", None):
            m = re.search(r"\((\d[\d,]*)\)", str(col(r, "Reviews")))
            count = int(m.group(1).replace(",", "")) if m else 0
        if rating in ("", None):
            m = re.match(r"\s*([\d.]+)", str(col(r, "Reviews")))
            rating = float(m.group(1)) if m else None
        city, st, zipc = parse_addr(str(col(r, "Address")))
        recs.append({
            "Name": clean_name(col(r, "Name")), "Listing name": str(col(r, "Name")),
            "Website": site, "Domain": norm_domain(site),
            "City": str(col(r, "City")) or city, "State": (str(col(r, "State")) or st).upper(),
            "Zip": str(col(r, "Zip")) or zipc, "County": str(col(r, "County")),
            "Phone": str(col(r, "Phone")), "Rating": rating, "Review count": int(count or 0),
            "Primary type": str(col(r, "Primary type")), "Source list": str(col(r, "Source list")),
            "Place ID": str(col(r, "Place ID")), "Google Maps URL": str(col(r, "Google Maps URL")),
        })
    return recs


def group_by_domain(recs):
    groups = {}
    for r in recs:
        if r["Domain"]:
            groups.setdefault(r["Domain"], []).append(r)
    companies = {}
    for dom, listings in groups.items():
        rep = dict(max(listings, key=lambda x: x["Review count"]))
        rep["Listings"] = len(listings)
        rep["DMV"] = "yes" if rep["State"] in DMV_STATES else "no"
        companies[dom] = rep
    no_site = [r for r in recs if not r["Domain"]]
    return companies, no_site


# ---------------------------------------------------------------- site cache

lock = threading.Lock()


def cache_path(run_date):
    return OUTREACH / f"site-cache-{run_date}.jsonl"


def load_site_cache(path):
    sites, llm = {}, {}
    if not path.exists():
        return sites, llm
    bad = 0
    # split on "\n" only: page text can contain   and other separators that str.splitlines() would break on
    for line in path.read_text(encoding="utf-8").split("\n"):
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except ValueError:
            bad += 1
            continue
        (llm if row.get("kind") == "llm" else sites)[row["domain"]] = row
    if bad:
        print(f"  site cache: skipped {bad} unreadable line(s); those domains will be fetched again", flush=True)
    return sites, llm


def append_cache(fh, row):
    with lock:
        fh.write(json.dumps(row, ensure_ascii=False) + "\n")
        fh.flush()


# ---------------------------------------------------------------- fetch and extract

def detect_platform(html, headers, listed_url):
    low = html.lower()
    hdr_blob = "\n".join(f"{k.lower()}: {str(v).lower()}" for k, v in (headers or {}).items())
    evidence = []
    if "utm_medium=wix_google_business_profile" in (listed_url or "").lower():
        evidence.append("listing URL carries wix_google_business_profile")
        return "wix", evidence
    for name, html_res, hdr_res in PLATFORM_RULES:
        for pat in html_res:
            if re.search(pat.lower(), low):
                evidence.append(f"html:{pat}")
                return name, evidence
        for pat in hdr_res:
            if re.search(pat, hdr_blob, re.M):
                evidence.append(f"header:{pat}")
                return name, evidence
    gen = re.search(r'<meta[^>]+name="generator"[^>]+content="([^"]+)"', html, re.I)
    if gen:
        evidence.append(f"generator:{gen.group(1)[:60]}")
    server = (headers or {}).get("Server") or (headers or {}).get("server")
    if server:
        evidence.append(f"server:{server[:40]}")
    return "custom", evidence


def extract(html, final_url):
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "noscript", "svg", "template", "iframe"]):
        t.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    meta = soup.find("meta", attrs={"name": re.compile("^description$", re.I)})
    desc = (meta.get("content") or "") if meta else ""
    heads = [h.get_text(" ", strip=True) for h in soup.find_all(["h1", "h2"])][:14]
    host = urlparse(final_url).netloc.lower().replace("www.", "")
    nav, candidates, seen = [], [], set()
    for a in soup.find_all("a", href=True):
        href = urljoin(final_url, a["href"].strip())
        p = urlparse(href)
        if p.scheme not in ("http", "https") or p.netloc.lower().replace("www.", "") != host:
            continue
        text = a.get_text(" ", strip=True)[:60]
        key = p.path.rstrip("/").lower()
        if not key or key in seen or any(key.endswith(ext) for ext in (".pdf", ".jpg", ".png", ".zip")):
            continue
        seen.add(key)
        if len(nav) < 40 and text:
            nav.append(text)
        if SUBPAGE_RE.search(key) or SUBPAGE_RE.search(text):
            candidates.append(href)
    text = re.sub(r"\s+", " ", soup.get_text(" ")).strip()
    return {"title": title[:200], "desc": desc[:300], "headings": heads, "nav": nav, "text": text,
            "subpage_candidates": candidates[:8]}


def scan_domain(domain, listed_url):
    rec = {"kind": "site", "domain": domain, "listed_url": listed_url, "fetched": time.strftime("%Y-%m-%d %H:%M")}
    r, err = fetch(listed_url or domain)
    if r is None:
        rec.update({"ok": False, "error": err})
        return rec
    try:
        html = r.text
        platform, evidence = detect_platform(html, r.headers, listed_url)
        home = extract(html, r.url)
        rec.update({
            "ok": True, "final_url": r.url, "https_ok": r.url.lower().startswith("https://"),
            "platform": platform, "platform_evidence": evidence[:4],
            "agency": "one_firefly" if ONEFIREFLY_RE.search(html.lower()) else "none",
            "title": home["title"], "desc": home["desc"], "headings": home["headings"], "nav": home["nav"],
            "text": home["text"][:HOME_TEXT_CAP], "text_len": len(home["text"]), "subpages": [],
        })
        for url in home["subpage_candidates"]:
            if len(rec["subpages"]) >= MAX_SUBPAGES:
                break
            if urlparse(url).path.rstrip("/") == urlparse(r.url).path.rstrip("/"):
                continue
            sr, _ = fetch(url, timeout=(5, 12))
            if sr is None:
                continue
            sub = extract(sr.text, sr.url)
            if len(sub["text"]) < 200:
                continue
            rec["subpages"].append({"url": sr.url, "title": sub["title"], "text": sub["text"][:SUB_TEXT_CAP]})
    except Exception as e:  # noqa: BLE001
        rec.update({"ok": False, "error": "parse:" + type(e).__name__})
    return rec


def fetch_all(companies, run_date, limit=0):
    path = cache_path(run_date)
    sites, _ = load_site_cache(path)
    targets = [(d, c["Website"]) for d, c in companies.items() if d not in sites]
    if limit:
        targets = targets[:limit]
    print(f"{len(companies)} companies, {len(sites)} cached, {len(targets)} to fetch", flush=True)
    n = 0
    with ThreadPoolExecutor(max_workers=THREADS) as ex, path.open("a", encoding="utf-8") as fh:
        futs = {ex.submit(scan_domain, d, u): d for d, u in targets}
        for fut in as_completed(futs):
            rec = fut.result()
            append_cache(fh, rec)
            sites[rec["domain"]] = rec
            n += 1
            if n % 100 == 0:
                print(f"  fetched {n}/{len(targets)}", flush=True)
    return sites


# ---------------------------------------------------------------- Opus 5

def user_message(company, site):
    parts = [
        f"Company listing: {company['Name']} ({company['City']}, {company['State']}). Google primary type: "
        f"{company.get('Primary type') or 'n/a'}. Rating {company.get('Rating')} on {company.get('Review count')} reviews.",
        f"Listed website: {company['Website']}  Final URL: {site.get('final_url')}",
        f"Detected platform (deterministic): {PLATFORM_LABELS.get(site.get('platform'), site.get('platform'))}; "
        f"agency: {site.get('agency')}",
        f"Title: {site.get('title')}", f"Meta description: {site.get('desc')}",
        "Headings: " + " | ".join(site.get("headings") or []),
        "Navigation: " + " | ".join(site.get("nav") or []),
        "Homepage text:\n" + (site.get("text") or ""),
    ]
    for sp in site.get("subpages") or []:
        parts.append(f"Subpage {sp['url']} ({sp.get('title')}):\n{sp['text']}")
    parts.append("Return the JSON object now.")
    return "\n\n".join(parts)


def request_params(system_prompt, content):
    return {
        "model": MODEL, "max_tokens": MAX_TOKENS,
        "system": [{"type": "text", "text": system_prompt, "cache_control": {"type": "ephemeral"}}],
        "messages": [{"role": "user", "content": content}],
        "output_config": {"effort": "low", "format": {"type": "json_schema", "schema": SCHEMA}},
    }


def parse_message(msg):
    if msg.stop_reason == "refusal":
        return None, "refusal:" + str(getattr(getattr(msg, "stop_details", None), "category", "") or "")
    if msg.stop_reason == "max_tokens":
        return None, "max_tokens"
    text = next((b.text for b in msg.content if b.type == "text"), "")
    try:
        return json.loads(text), None
    except ValueError:
        return None, "bad_json"


def usage_cost(u, batch):
    in_rate, out_rate = (BATCH_IN_PER_M, BATCH_OUT_PER_M) if batch else (SYNC_IN_PER_M, SYNC_OUT_PER_M)
    cache_read = getattr(u, "cache_read_input_tokens", 0) or 0
    cache_write = getattr(u, "cache_creation_input_tokens", 0) or 0
    cache_read_rate = CACHE_READ_PER_M / (2 if batch else 1)
    cost = (u.input_tokens * in_rate + cache_write * in_rate * 1.25 + cache_read * cache_read_rate
            + u.output_tokens * out_rate) / 1e6
    return round(cost, 6)


def llm_record(domain, data, err, msg=None, batch=False, kind="llm"):
    rec = {"kind": kind, "domain": domain, "model": MODEL, "ts": time.strftime("%Y-%m-%d %H:%M"),
           "ok": data is not None, "error": err, "data": data}
    if msg is not None:
        u = msg.usage
        rec["usage"] = {"in": u.input_tokens, "out": u.output_tokens,
                        "cache_read": getattr(u, "cache_read_input_tokens", 0) or 0,
                        "cache_write": getattr(u, "cache_creation_input_tokens", 0) or 0}
        rec["cost_usd"] = usage_cost(u, batch)
    return rec


def estimate_cost(system_prompt, contents, batch):
    in_tokens = sum(len(c) for c in contents) / 4 + len(system_prompt) / 4  # system cached after the first
    out_tokens = EST_OUTPUT_TOKENS * len(contents)
    in_rate, out_rate = (BATCH_IN_PER_M, BATCH_OUT_PER_M) if batch else (SYNC_IN_PER_M, SYNC_OUT_PER_M)
    return round((in_tokens * in_rate + out_tokens * out_rate) / 1e6, 2)


def batches_path(run_date):
    return OUTREACH / f"qualify-batches-{run_date}.json"


def run_sync(client, todo, system_prompt, run_date, meter):
    path = cache_path(run_date)
    results = {}

    def one(domain, content):
        try:
            msg = client.messages.create(**request_params(system_prompt, content))
        except Exception as e:  # noqa: BLE001
            return llm_record(domain, None, "api:" + type(e).__name__ + ":" + str(e)[:160])
        data, err = parse_message(msg)
        return llm_record(domain, data, err, msg)

    with ThreadPoolExecutor(max_workers=4) as ex, path.open("a", encoding="utf-8") as fh:
        futs = {ex.submit(one, d, c): d for d, c in todo}
        for fut in as_completed(futs):
            rec = fut.result()
            append_cache(fh, rec)
            results[rec["domain"]] = rec
            meter.add("opus_sync", rec.get("cost_usd", 0))
    return results


def submit_batches(client, todo, system_prompt, run_date, params_fn=None, state_path=None):
    """todo: [(key, content)]. params_fn(key, content) builds the request; the default is this script's own read.
    state_path: another script's batch state file (hero_check.py keeps its own). Both default to this script's."""
    params_fn = params_fn or (lambda key, content: request_params(system_prompt, content))
    state_path = state_path or batches_path(run_date)
    state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {"batches": []}
    for i in range(0, len(todo), BATCH_CHUNK):
        chunk = todo[i:i + BATCH_CHUNK]
        ids = {f"d{i + k:06d}": d for k, (d, _) in enumerate(chunk)}
        reqs = [{"custom_id": cid, "params": params_fn(d, c)}
                for cid, (d, c) in zip(ids, chunk)]
        b = client.messages.batches.create(requests=reqs)
        state["batches"].append({"id": b.id, "ids": ids, "status": b.processing_status, "collected": False,
                                 "submitted": time.strftime("%Y-%m-%d %H:%M")})
        state_path.write_text(json.dumps(state, indent=1), encoding="utf-8")
        print(f"  submitted batch {b.id} with {len(reqs)} requests", flush=True)
    return state


def collect_batches(client, run_date, meter, wait=True, state_path=None, on_result=None):
    """on_result(key, message_or_None, error_or_None) records one answer and returns the record (with cost_usd).
    The default writes this script's llm record into the site cache. Returns {key: record}."""
    state_path = state_path or batches_path(run_date)
    if not state_path.exists():
        return {}
    state = json.loads(state_path.read_text(encoding="utf-8"))
    path = cache_path(run_date)
    results = {}

    def default_on_result(domain, msg, err):
        with path.open("a", encoding="utf-8") as fh:
            if msg is not None:
                data, perr = parse_message(msg)
                rec = llm_record(domain, data, perr, msg, batch=True)
            else:
                rec = llm_record(domain, None, err)
            append_cache(fh, rec)
        return rec

    on_result = on_result or default_on_result
    pending = [b for b in state["batches"] if not b.get("collected")]
    while pending:
        for b in pending:
            info = client.messages.batches.retrieve(b["id"])
            b["status"] = info.processing_status
            if info.processing_status != "ended":
                c = info.request_counts
                print(f"  {b['id']}: {info.processing_status}, processing {c.processing}, "
                      f"succeeded {c.succeeded}, errored {c.errored}", flush=True)
                continue
            for res in client.messages.batches.results(b["id"]):
                domain = b["ids"].get(res.custom_id)
                if not domain:
                    continue
                if res.result.type == "succeeded":
                    rec = on_result(domain, res.result.message, None)
                else:
                    etype = getattr(getattr(res.result, "error", None), "type", res.result.type)
                    rec = on_result(domain, None, f"batch:{res.result.type}:{etype}")
                results[domain] = rec
                meter.add("opus_batch", (rec or {}).get("cost_usd", 0))
            b["collected"] = True
            print(f"  collected {b['id']}: {len(b['ids'])} results", flush=True)
        state_path.write_text(json.dumps(state, indent=1), encoding="utf-8")
        pending = [b for b in state["batches"] if not b.get("collected")]
        if pending and not wait:
            print(f"{len(pending)} batch(es) still processing. Rerun later to collect.", flush=True)
            break
        if pending:
            time.sleep(POLL_SECONDS)
    return results


# ---------------------------------------------------------------- scoring

def is_thin(site):
    return bool(site) and site.get("ok") and site.get("text_len", 0) < THIN_TEXT and not site.get("subpages")


def score(company, site, llm):
    """Returns (tier, reason). tier in {A, B, C, exclude, unverified}."""
    if not site or not site.get("ok"):
        return "unverified", f"site not fetched: {(site or {}).get('error', 'not scanned')}"
    if site.get("agency") == "one_firefly":
        return "exclude", "One Firefly site (already agencied)"
    if is_thin(site):
        return "unverified", "thin page (JS-rendered or blocked), not read by the model"
    if not llm or not llm.get("ok"):
        return "unverified", f"model read failed: {(llm or {}).get('error', 'not run')}"
    d = llm["data"]
    if not d["is_integrator"]:
        return "exclude", "not an integrator: " + d["reason"]
    if d["market_focus"] in ("commercial", "mixed_commercial_lead"):
        return "exclude", f"{d['market_focus'].replace('_', ' ')}: " + d["reason"]
    off = d["offerings"]
    core = sum([off["lighting_control"], off["automation_control"], off["audio_video"] or off["home_theater"]])
    weak = site.get("platform") in WEAK_PLATFORMS
    conf = float(d.get("confidence") or 0)
    if d["market_focus"] == "unclear" or conf < 0.6 or core < 2:
        return "C", f"review: focus {d['market_focus']}, {core} core offerings, confidence {conf:.2f}"
    if weak:
        return "A", f"residential integrator on {PLATFORM_LABELS.get(site.get('platform'))}, {core} core offerings"
    return "B", f"residential integrator on {PLATFORM_LABELS.get(site.get('platform'))}, {core} core offerings"


SCORED_COLS = ["Tier", "Name", "Website", "Domain", "City", "State", "Zip", "DMV", "Phone", "Rating", "Review count",
               "Listings", "Platform", "Weak hosting", "Agency", "Market focus", "Is integrator", "Core offerings",
               "Offerings", "Brands", "High-end signals", "Service pages", "Area pages", "Gallery", "Testimonials",
               "Single page", "Confidence", "Model reason", "Score reason", "Site title", "Final URL", "HTTPS",
               "Source list", "Place ID", "Google Maps URL"]


def build_rows(companies, sites, llms):
    rows = []
    for dom, c in companies.items():
        site, llm = sites.get(dom), llms.get(dom)
        tier, why = score(c, site, llm)
        d = (llm or {}).get("data") or {}
        off = d.get("offerings") or {}
        sq = d.get("site_quality") or {}
        rows.append({
            "Tier": tier, "Name": c["Name"], "Website": c["Website"], "Domain": dom, "City": c["City"],
            "State": c["State"], "Zip": c["Zip"], "DMV": c["DMV"], "Phone": c["Phone"], "Rating": c["Rating"],
            "Review count": c["Review count"], "Listings": c["Listings"],
            "Platform": PLATFORM_LABELS.get((site or {}).get("platform"), "") if site else "",
            "Weak hosting": "yes" if (site or {}).get("platform") in WEAK_PLATFORMS else "no",
            "Agency": (site or {}).get("agency", ""), "Market focus": d.get("market_focus", ""),
            "Is integrator": ("yes" if d.get("is_integrator") else "no") if d else "",
            "Core offerings": (sum([off.get("lighting_control", False), off.get("automation_control", False),
                                    off.get("audio_video", False) or off.get("home_theater", False)]) if off else ""),
            "Offerings": ", ".join(k for k, v in off.items() if v),
            "Brands": ", ".join(d.get("brands") or []), "High-end signals": "; ".join(d.get("high_end_signals") or []),
            "Service pages": _yn(sq.get("has_service_pages")), "Area pages": _yn(sq.get("has_service_area_pages")),
            "Gallery": _yn(sq.get("has_gallery")), "Testimonials": _yn(sq.get("has_testimonials")),
            "Single page": _yn(sq.get("single_page")), "Confidence": d.get("confidence", ""),
            "Model reason": d.get("reason", ""), "Score reason": why, "Site title": (site or {}).get("title", ""),
            "Final URL": (site or {}).get("final_url", ""), "HTTPS": _yn((site or {}).get("https_ok")),
            "Source list": c.get("Source list", ""), "Place ID": c.get("Place ID", ""),
            "Google Maps URL": c.get("Google Maps URL", ""),
        })
    order = {"A": 0, "B": 1, "C": 2, "unverified": 3, "exclude": 4}
    rows.sort(key=lambda r: (order[r["Tier"]], r["State"], r["Name"]))
    return rows


def _yn(v):
    return "" if v is None else ("yes" if v else "no")


def write_outputs(rows, run_date, meter, no_site_count):
    out = OUTREACH / f"qualified-seed-{run_date}.xlsx"
    wb = openpyxl.Workbook()
    sheets = [("Qualified", [r for r in rows if r["Tier"] in ("A", "B", "C")]),
              ("Excluded", [r for r in rows if r["Tier"] == "exclude"]),
              ("Unverified", [r for r in rows if r["Tier"] == "unverified"]),
              ("All scored", rows)]
    first = True
    for title, subset in sheets:
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        first = False
        ws.append(SCORED_COLS)
        for r in subset:
            ws.append([r.get(c, "") if r.get(c) is not None else "" for c in SCORED_COLS])
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
    ws = wb.create_sheet("Method")
    for line in [
        f"Built {run_date} by scripts/qualify_sites.py. Model: {MODEL}, structured JSON, effort low.",
        "One company per website domain; the highest-reviewed listing represents it.",
        "Fetch: homepage plus up to two same-domain pages (services, residential, about, solutions).",
        "Platform: deterministic fingerprints (Wix, Squarespace, GoDaddy, Shopify, Weebly, Duda, Webflow, HubSpot, "
        "WordPress, else custom/unknown). Weak hosting = WordPress, Wix, Squarespace, GoDaddy, Shopify, Weebly, Duda "
        "(Jonathan's list, 2026-09-13).",
        "Exclude: One Firefly in the HTML; model says not an integrator; commercial or mixed commercial lead.",
        "Tier A: residential (or mixed, residential lead) integrator, 2+ of lighting control / automation / AV or "
        "theater, weak hosting, confidence >= 0.6.",
        "Tier B: same as A on a custom or unknown platform.",
        "Tier C: residential integrator with one core offering, unclear focus, or confidence < 0.6. Review by hand.",
        "Unverified: site could not be fetched or the model read failed. Not in the LinkedIn CSVs.",
        "DMV = yes marks DC, MD, VA. Kept in linkedin-company-list-<DATE>.csv (decision 2026-09-13); removed in the "
        "-no-dmv variant. The 0 to 30 review rule is not applied to the ad seed; Review count is a column.",
        f"Seed rows with no website skipped: {no_site_count}.",
        f"Spend this run: {json.dumps(meter.summary())}",
        "LinkedIn Campaign Manager company list: min 300 rows, max 300,000, 20 MB, header must match the template "
        "downloaded from Campaign Manager. Verify the header before uploading.",
    ]:
        ws.append([line])
    wb.save(out)

    ab = [r for r in rows if r["Tier"] in ("A", "B")]
    csv_all = OUTREACH / f"linkedin-company-list-{run_date}.csv"
    csv_no_dmv = OUTREACH / f"linkedin-company-list-{run_date}-no-dmv.csv"
    n_all = write_linkedin_csv(ab, csv_all)
    n_no_dmv = write_linkedin_csv([r for r in ab if r["DMV"] == "no"], csv_no_dmv)
    if n_all < 300:
        print(f"WARNING: only {n_all} Tier A+B companies; LinkedIn needs at least 300 rows to upload.", flush=True)
    return out, csv_all, csv_no_dmv, n_all, n_no_dmv


def summarize(rows, meter, out, csv_all, csv_no_dmv, n_all, n_no_dmv):
    def count(key):
        c = {}
        for r in rows:
            c[r[key] or "?"] = c.get(r[key] or "?", 0) + 1
        return dict(sorted(c.items(), key=lambda kv: -kv[1]))
    ab = [r for r in rows if r["Tier"] in ("A", "B")]
    return {
        "companies": len(rows), "by_tier": count("Tier"), "by_platform": count("Platform"),
        "tier_ab_by_state": dict(sorted(
            {s: sum(1 for r in ab if r["State"] == s) for s in {r["State"] for r in ab}}.items(),
            key=lambda kv: -kv[1])[:15]),
        "linkedin_rows": n_all, "linkedin_rows_no_dmv": n_no_dmv, "cost": meter.summary(),
        "xlsx": str(out), "csv": str(csv_all), "csv_no_dmv": str(csv_no_dmv),
    }


# ---------------------------------------------------------------- main

def places_spend_for(run_date):
    p = OUTREACH / f"places-cache-{run_date}.jsonl"
    if not p.exists():
        return 0.0
    n = sum(1 for line in p.read_text(encoding="utf-8").splitlines() if line.strip())
    return round(n * PLACES_PRICE_PER_REQUEST, 4)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--input", default=None, help="seed workbook (default: newest projects/outreach/places-seed-*.xlsx)")
    ap.add_argument("--limit", type=int, default=0, help="only the first N unfetched domains (smoke test)")
    ap.add_argument("--dry-run", action="store_true", help="fetch and fingerprint, skip the model, print estimate")
    ap.add_argument("--no-batch", action="store_true", help="synchronous calls at full price (small runs)")
    ap.add_argument("--no-wait", action="store_true", help="submit the batch and exit without polling")
    ap.add_argument("--build-only", action="store_true", help="score and write outputs from the cache only")
    ap.add_argument("--max-cost", type=float, default=None, help="USD cap for the whole run (default SEED_BUDGET_USD)")
    ap.add_argument("--date", default=date.today().isoformat())
    a = ap.parse_args(argv)

    seed = Path(a.input) if a.input else newest("places-seed-*.xlsx")
    if not seed or not seed.exists():
        sys.exit("No seed workbook. Run scripts/places_seed.py first or pass --input.")
    recs = load_seed(seed)
    companies, no_site = group_by_domain(recs)
    print(f"seed {seed.name}: {len(recs)} rows, {len(companies)} companies with a website, {len(no_site)} without",
          flush=True)

    already = {"places_requests": places_spend_for(a.date)}
    _, llm_cached = load_site_cache(cache_path(a.date))
    for rec in llm_cached.values():
        already["opus_prior"] = round(already.get("opus_prior", 0) + rec.get("cost_usd", 0), 4)
    meter = CostMeter(budget_cap(override=a.max_cost), already)

    if a.build_only:
        sites, llms = load_site_cache(cache_path(a.date))
    else:
        sites = fetch_all(companies, a.date, a.limit)
        _, llms = load_site_cache(cache_path(a.date))
        system_prompt = build_system_prompt()
        todo = [(d, user_message(companies[d], sites[d])) for d in companies
                if d in sites and sites[d].get("ok") and sites[d].get("agency") != "one_firefly"
                and not is_thin(sites[d]) and not (llms.get(d) or {}).get("ok")]
        if a.limit:
            todo = todo[:a.limit]
        est = estimate_cost(system_prompt, [c for _, c in todo], batch=not a.no_batch)
        print(f"{len(todo)} sites for Opus 5 ({'sync' if a.no_batch else 'batch'}), estimated ${est:.2f}; "
              f"spent so far ${meter.total:.2f} of ${meter.cap:.2f}", flush=True)
        if a.dry_run:
            rows = build_rows(companies, sites, llms)
            plat = {}
            for r in rows:
                plat[r["Platform"] or "(not fetched)"] = plat.get(r["Platform"] or "(not fetched)", 0) + 1
            print(json.dumps({"platforms": plat, "opus_estimate_usd": est, "cost": meter.summary()}, indent=2))
            return 0
        if todo:
            try:
                meter.assert_affordable(est, "Opus 5 qualification")
            except BudgetExceeded as e:
                sys.exit(f"STOPPED: {e}")
            import anthropic
            client = anthropic.Anthropic(api_key=get_secret(
                "ANTHROPIC_API_KEY", required=True,
                hint="Copy it from the Monarc OS main machine into %USERPROFILE%\\.monarc\\secrets.env."))
            if a.no_batch:
                llms.update(run_sync(client, todo, system_prompt, a.date, meter))
            else:
                pending = batches_path(a.date).exists() and any(
                    not b.get("collected") for b in json.loads(batches_path(a.date).read_text(encoding="utf-8"))["batches"])
                if not pending:
                    submit_batches(client, todo, system_prompt, a.date)
                else:
                    print("Pending batch found; collecting it before submitting anything new.", flush=True)
                llms.update(collect_batches(client, a.date, meter, wait=not a.no_wait))
        elif batches_path(a.date).exists():
            import anthropic
            client = anthropic.Anthropic(api_key=get_secret("ANTHROPIC_API_KEY", required=True))
            llms.update(collect_batches(client, a.date, meter, wait=not a.no_wait))

    rows = build_rows(companies, sites, llms)
    out, csv_all, csv_no_dmv, n_all, n_no_dmv = write_outputs(rows, a.date, meter, len(no_site))
    print(json.dumps(summarize(rows, meter, out, csv_all, csv_no_dmv, n_all, n_no_dmv), indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
