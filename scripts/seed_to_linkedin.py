"""Write a places_seed workbook as a LinkedIn company-list CSV (Campaign Manager's Account Match template:
companyname, companywebsite, companyemaildomain, linkedincompanypageurl, stocksymbol, industry, city, state,
companycountry, zipcode). One row per website domain; the listing with the most reviews names the company; rows
without a website are skipped (LinkedIn matches on the domain); rows merged in from other lists (Source list not
"places-...") are skipped unless --all.

Usage:
  python scripts/seed_to_linkedin.py projects/outreach/places-seed-2026-09-13.xlsx linkedin-company-list-2026-09-24-integrators-full.csv
  python scripts/seed_to_linkedin.py <seed.xlsx> <out.csv> --all      # keep the merged-in rows too

Patrick's template, 2026-09-24 (the same header the 2026-09-13 upload used).
"""
import argparse
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from outreach_common import OUTREACH, STATE_NAMES, norm_domain, write_linkedin_csv  # noqa: E402


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("seed")
    ap.add_argument("out")
    ap.add_argument("--all", action="store_true", help="keep rows merged in from other lists")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--linkedin-from", action="append", default=[],
                    help="a meta-seed workbook (Seed sheet: Domain, LinkedIn) whose company page links are merged in by domain; repeatable")
    ap.add_argument("--industry", default="", help="LinkedIn industry label written on every row (their taxonomy), blank to leave it out")
    a = ap.parse_args(argv)

    pages = {}
    for path in a.linkedin_from:
        wb2 = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws2 = wb2["Seed"] if "Seed" in wb2.sheetnames else wb2.worksheets[0]
        it2 = ws2.iter_rows(values_only=True)
        h2 = [str(h or "").strip() for h in next(it2)]
        try:
            i_dom, i_li = h2.index("Domain"), h2.index("LinkedIn")
        except ValueError:
            sys.exit(f"{path}: needs Domain and LinkedIn columns")
        import re
        for r in it2:
            if r and r[i_dom] and r[i_li] and "/company/" in str(r[i_li]):
                # the crawler keeps whatever the site linked (/about, /posts, /admin ...); Campaign Manager wants
                # the bare page address (2026-09-24, an upload was refused)
                m = re.search(r"linkedin\.com/company/([^/?#\s]+)", str(r[i_li]), re.I)
                if m:
                    pages.setdefault(norm_domain(r[i_dom]) or str(r[i_dom]).lower(), f"https://www.linkedin.com/company/{m.group(1)}")

    wb = openpyxl.load_workbook(a.seed, read_only=True, data_only=True)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = [str(h or "").strip() for h in next(it)]
    ix = {h: i for i, h in enumerate(head)}

    def col(r, name):
        i = ix.get(name)
        return r[i] if i is not None and i < len(r) else None

    best = {}
    skipped = {"no_website": 0, "merged_in": 0, "not_us": 0}
    for r in it:
        if not r or not col(r, "Name"):
            continue
        src = str(col(r, "Source list") or "")
        if not a.all and src and not src.startswith("places-"):
            skipped["merged_in"] += 1
            continue
        dom = norm_domain(col(r, "Website"))
        if not dom:
            skipped["no_website"] += 1
            continue
        if str(col(r, "State") or "").upper() not in STATE_NAMES:  # a few foreign listings ride along; LinkedIn wants a US state
            skipped["not_us"] += 1
            continue
        reviews = int(col(r, "Review count") or 0)
        cur = best.get(dom)
        if cur is None or reviews > cur["_reviews"]:
            best[dom] = {"Name": str(col(r, "Name")).strip(), "Website": col(r, "Website") or "", "Domain": dom,
                         "City": col(r, "City") or "", "State": str(col(r, "State") or "").upper(),
                         "Zip": col(r, "Zip") or "", "_reviews": reviews,
                         "LinkedIn": pages.get(dom, ""), "Industry": a.industry}
    rows = sorted(best.values(), key=lambda x: (x["State"], x["City"], x["Name"]))
    out = Path(a.out)
    if not out.is_absolute() and out.parent == Path("."):
        out = OUTREACH / out
    n = write_linkedin_csv(rows, out)
    with_page = sum(1 for r in rows if r["LinkedIn"])
    print(f"{out}: {n} companies (one per domain), {with_page} with a LinkedIn page; skipped {skipped}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
