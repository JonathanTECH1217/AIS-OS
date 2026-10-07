"""Give every company in the Prospects call list a Vertical (Jonathan, 2026-10-01: "I want all of the companies in
the Airtable list to be given a vertical that they cover").

Usage: python scripts/label_verticals.py               dry run: counts, samples, the review CSV; writes nothing
       python scripts/label_verticals.py --ai          also runs the AI read for the sites read as "not an integrator"
       python scripts/label_verticals.py --write       writes the changed labels to Airtable (after a dry run)
       add --cached to reuse the last pull of the table instead of pulling it again (dry runs only)
       add --limit N with --ai to read only the first N (a smoke test)
       add --crawl (with --ai) to read the homepages of the rows that would fall to the list default

The twelve choices (his five, then the ones common in the list; his pick of the "full set"): Electrician,
AV integrator, HVAC, Plumber, Roofer, Security, Low voltage / IT, Generator, Automotive, General contractor,
Multi-trade (two or more of plumbing, HVAC, electrical), Not a contractor (stores, makers, venues, agencies).

How each row is labeled, first rule that applies:
  1. The integrators he checked by hand (Queue 1 to 867): AV integrator.
  2. A site read from 2026-09-13 (projects/outreach/qualified-seed-2026-09-13.xlsx, scripts/qualify_sites.py):
     "is an integrator" -> AV integrator; "not an integrator" -> the AI picks one of the twelve from the read's own
     one-line reason (claude-opus-5, effort low, cached in vertical-ai-<date>.jsonl, so a rerun never pays twice).
  3. Trade words in the company name and web address (two of plumbing, HVAC, electrical = Multi-trade).
  4. Google's own category for the listing (the Places sweeps' "Primary type").
  5. With --crawl: the rows still unsure get their homepage's title, description, and first heading read (one GET,
     cached in vertical-sites-<date>.jsonl), and the same AI read picks from those.
  6. The list it came from: the integrator rows AV integrator, the electrician list Electrician.
Output: projects/outreach/vertical-labels-<date>.csv (every row, old and new label, which rule decided).
--write keeps the old labels in vertical-before-<date>.json first, then patches only the rows that change, ten at a
time with typecast, which also adds any new choice to the column. Airtable's API cannot move a column; the Vertical
column is dragged next to Name in the grid by hand. Never prints a key.
"""
import collections
import csv
import json
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl
import requests

sys.path.insert(0, str(Path(__file__).resolve().parent))
from outreach_common import ROOT, load_env, norm_domain  # noqa: E402

OUT = ROOT / "projects" / "outreach"
DATE = "2026-10-01"
BASE, CONTACTS = "appPM1HmRfFlgA484", "tblGdourIuqa9onmp"
API = "https://api.airtable.com/v0"
CACHE = OUT / f"vertical-pull-{DATE}.json"
AI_CACHE = OUT / f"vertical-ai-{DATE}.jsonl"
CHECKED_LAST_QUEUE = 867  # 842 integrators checked by hand + 25 checked ones turned down in September
MODEL = "claude-opus-5"
IN_PER_M, OUT_PER_M = 5.0, 25.0
BATCH = 40

LABELS = ["Electrician", "AV integrator", "HVAC", "Plumber", "Roofer", "Security", "Low voltage / IT", "Generator",
          "Automotive", "General contractor", "Multi-trade", "Not a contractor"]

# Trade words in the name and web address, checked in this order (a car audio shop's "audio" is not AV; an
# "Electric & Generator" shop is a generator dealer; "Audio Video & Security" is an integrator)
P = re.compile(r"plumb|rooter|\bdrains?\b|sewer|septic|water heater", re.I)
H = re.compile(r"\bhvac\b|heating|\bcooling\b|air condition|furnace|refrigerat|\bmechanical\b|"
               r"(heat|heating|cool|cooling)\W+(and|&|n)?\W*air\b|\bair\W+(and|&)\W*(heat|heating|cool)", re.I)
E = re.compile(r"electric|electrician|\bsparky\b|\belec\b", re.I)
RULES = [
    ("Not a contractor", re.compile(r"charging station|electric motor", re.I)),  # public EV stations, motor repair shops
    ("Electrician", re.compile(r"electric vehicle|\bev charg", re.I)),            # EV charger installers
    ("Automotive", re.compile(r"\bauto\b(?!mation)|automotive|\bcars?\b|car audio|\btrucks?\b|vehicle|"
                              r"\btint|remote start|diesel|\btires?\b|collision|\bmarine\b|\brv\b|mobile electronics", re.I)),
    ("Roofer", re.compile(r"\broof|gutter", re.I)),
    ("Generator", re.compile(r"generator", re.I)),
    ("Plumber", P),
    ("HVAC", H),
    ("AV integrator", re.compile(r"audio|video|\ba/?v\b|audiovisual|theat(er|re)|cinema|home automation|smart ?home|"
                                 r"integrat(ion|ors?)|control4|crestron|savant|sonos|acoustic|\bsound\b|media room|"
                                 r"home tech|hi-?fi|stereo|lutron|\bavi\b", re.I)),
    ("Security", re.compile(r"security|alarm|surveillance|cctv|\bcameras?\b|access control|locksmith", re.I)),
    ("Electrician", E),
    ("Electrician", re.compile(r"solar|photovolt", re.I)),  # solar installers do the electrical work
    ("Low voltage / IT", re.compile(r"low ?voltage|cabling|structured wiring|\bdata\b|networks?\b|fiber|telecom|"
                                    r"communications|computers?|wireless|satellite|antenna|\bdish\b|voip|\bIT\b")),
    ("General contractor", re.compile(r"construction|builders?\b|contracting|remodel|renovat|general contractor|"
                                      r"handyman|restoration", re.I)),
]
GOOGLE = {"plumber": "Plumber", "roofing_contractor": "Roofer", "electrician": "Electrician",
          "electric_vehicle_charging_station": "Electrician", "car_repair": "Automotive", "auto_parts_store": "Automotive",
          "truck_dealer": "Automotive", "car_dealer": "Automotive", "telecommunications_service_provider": "Low voltage / IT"}
NOT_CONTRACTOR_TYPES = {"movie_theater", "performing_arts_theater", "live_music_venue", "hotel", "travel_agency",
                        "amusement_center", "real_estate_agency", "television_studio", "health", "educational_institution",
                        "non_profit_organization", "government_office", "local_government_office",
                        "association_or_organization", "furniture_store", "home_goods_store", "sporting_goods_store",
                        "building_materials_store", "hardware_store", "wholesaler", "manufacturer"}

SYSTEM = f"""You label companies on a cold-call list by the trade they work in. For each company you get its name and either
a one-line description written by someone who read its website, or the site's own title, description, and first
heading. Pick exactly one label from this list:

- AV integrator: designs or installs audio, video, home theater, automation, lighting control, shades, or smart home
  systems in homes or businesses. Includes commercial AV integrators and TV mounting installers.
- Electrician: electrical contractor (wiring, panels, service upgrades, EV chargers, lighting installs, solar electrical).
- HVAC: heating, cooling, air conditioning, ventilation, or refrigeration contractor.
- Plumber: plumbing, drains, water heaters, sewer, septic.
- Roofer: roofing or gutters.
- Security: alarms, surveillance cameras, access control, fire alarm, or security monitoring.
- Low voltage / IT: structured cabling, networks, managed IT, telecom, phone systems, fiber, satellite or antennas.
- Generator: sells and installs standby generators.
- Automotive: car audio, car stereos, auto electric, window tint, remote start, marine or vehicle work of any kind.
- General contractor: builders, remodelers, handymen, restoration, renovation, design-build, construction services.
- Multi-trade: a home services company doing two or more of HVAC, plumbing, and electrical.
- Not a contractor: stores and retailers, manufacturers, distributors, wholesalers, manufacturers' reps, AV rental,
  staging, or event production, video or media production, venues and theaters, software, agencies, real estate,
  clinics, schools, nonprofits, industrial automation, warranty plans, and anything else that does not install or
  service systems in buildings.

- Unknown: the description says the site was blocked, empty, parked, expired, or gave no business information.

Read past the "no ..." part of each description: "electrical contractor ...; no AV work" is an Electrician. Use only
these exact labels: {", ".join(LABELS)}, Unknown. Return one entry per company, with its number."""
SCHEMA = {"type": "object", "additionalProperties": False, "required": ["labels"],
          "properties": {"labels": {"type": "array", "items": {
              "type": "object", "additionalProperties": False, "required": ["i", "vertical"],
              "properties": {"i": {"type": "integer"}, "vertical": {"type": "string", "enum": LABELS + ["Unknown"]}}}}}}

env, _ = load_env()


def digits(v):
    d = re.sub(r"\D", "", str(v or ""))
    return d[-10:] if len(d) >= 10 else ""


def airtable():
    s = requests.Session()
    s.headers["Authorization"] = f"Bearer {env['AIRTABLE_API_KEY']}"
    return s


def call(s, method, url, **kw):
    for _ in range(5):
        r = s.request(method, url, timeout=60, **kw)
        if r.status_code == 429:
            time.sleep(30)
            continue
        time.sleep(0.22)  # five requests a second at most
        return r
    return r


def pull(cached):
    if cached and CACHE.exists():
        return json.loads(CACHE.read_text(encoding="utf-8"))
    s, rows, offset = airtable(), [], None
    fields = ["Name", "Website", "Phone", "Vertical", "Queue", "Company ID", "Status (typed)"]
    while True:
        params = [("pageSize", 100)] + [("fields[]", f) for f in fields] + ([("offset", offset)] if offset else [])
        r = call(s, "GET", f"{API}/{BASE}/{CONTACTS}", params=params)
        r.raise_for_status()
        j = r.json()
        rows += [{"id": x["id"], **x["fields"]} for x in j["records"]]
        offset = j.get("offset")
        if not offset:
            break
    CACHE.write_text(json.dumps(rows), encoding="utf-8")
    return rows


def site_reads():
    ws = openpyxl.load_workbook(OUT / "qualified-seed-2026-09-13.xlsx", read_only=True, data_only=True)["All scored"]
    it = ws.iter_rows(values_only=True)
    head = [str(h or "") for h in next(it)]
    ix = {h: i for i, h in enumerate(head)}
    out = {}
    for r in it:
        d = norm_domain(str(r[ix["Domain"]] or r[ix["Website"]] or ""))
        if d and d not in out:
            out[d] = {"verdict": str(r[ix["Is integrator"]] or ""), "reason": str(r[ix["Model reason"]] or ""),
                      "focus": str(r[ix["Market focus"]] or ""), "title": str(r[ix["Site title"]] or "")}
    return out


def google_types():
    out = {}
    for name in ["places-seed-2026-09-24-electricians.xlsx", "places-seed-2026-09-13.xlsx"]:
        ws = openpyxl.load_workbook(OUT / name, read_only=True, data_only=True)["Integrators"]
        it = ws.iter_rows(values_only=True)
        head = [str(h or "") for h in next(it)]
        ix = {h: i for i, h in enumerate(head)}
        for r in it:
            t = str(r[ix["Primary type"]] or "")
            d, tel = norm_domain(str(r[ix["Website"]] or r[ix["Domain"]] or "")), digits(r[ix["Phone"]])
            if d:
                out.setdefault(("d", d), t)
            if tel:
                out.setdefault(("t", tel), t)
    return out


def by_words(name, website, extra=""):
    dom = norm_domain(str(website or ""))
    dom_words = re.sub(r"\.(com|net|org|us|co|biz|info|pro|io|tv|llc)(\.[a-z]{2})?$", "", dom).replace("-", " ").replace(".", " ")
    text = f"{name or ''} | {dom_words} | {extra or ''}"
    if sum(bool(p.search(text)) for p in (P, H, E)) >= 2:
        return "Multi-trade"
    return next((label for label, p in RULES if p.search(text)), None)


def ai_cache():
    done = {}
    if AI_CACHE.exists():
        for line in AI_CACHE.read_text(encoding="utf-8").splitlines():
            rec = json.loads(line)
            if rec.get("vertical") in LABELS + ["Unknown"]:  # Unknown is cached too (no second charge) but decides nothing
                done[rec["domain"]] = rec["vertical"]
    return done


def run_ai(todo):
    """todo: {domain: (name, focus, reason)}. Labels the ones not cached; returns the cost in dollars."""
    import anthropic
    client = anthropic.Anthropic(api_key=env["ANTHROPIC_API_KEY"])
    items = sorted(todo.items())
    chunks = [items[i:i + BATCH] for i in range(0, len(items), BATCH)]
    print(f"AI read: {len(items)} companies in {len(chunks)} requests ({MODEL}, effort low)")

    def one(chunk):
        lines = [f"{n}. {name} | {focus} | {reason}" for n, (_, (name, focus, reason)) in enumerate(chunk, 1)]
        try:
            msg = client.messages.create(
                model=MODEL, max_tokens=6000,
                system=[{"type": "text", "text": SYSTEM, "cache_control": {"type": "ephemeral"}}],
                messages=[{"role": "user", "content": "Companies (number. name | market focus | description):\n" + "\n".join(lines)}],
                output_config={"effort": "low", "format": {"type": "json_schema", "schema": SCHEMA}})
        except Exception as e:  # noqa: BLE001
            return [], 0.0, f"api:{type(e).__name__}:{str(e)[:120]}"
        u = msg.usage
        cost = (u.input_tokens * IN_PER_M + (getattr(u, "cache_creation_input_tokens", 0) or 0) * IN_PER_M * 1.25
                + (getattr(u, "cache_read_input_tokens", 0) or 0) * IN_PER_M * 0.1 + u.output_tokens * OUT_PER_M) / 1e6
        text = next((b.text for b in msg.content if b.type == "text"), "")
        try:
            got = {e["i"]: e["vertical"] for e in json.loads(text)["labels"]}
        except (ValueError, KeyError, TypeError):
            return [], cost, "bad_json"
        return [(dom, got.get(n)) for n, (dom, _) in enumerate(chunk, 1)], cost, None

    total, missing = 0.0, 0
    with ThreadPoolExecutor(max_workers=4) as ex, AI_CACHE.open("a", encoding="utf-8") as fh:
        for k, (pairs, cost, err) in enumerate(ex.map(one, chunks), 1):
            total += cost
            if err:
                print(f"  request {k}: {err}")
            for dom, label in pairs:
                if label in LABELS + ["Unknown"]:
                    fh.write(json.dumps({"domain": dom, "vertical": label, "model": MODEL}) + "\n")
                else:
                    missing += 1
            if k % 10 == 0:
                print(f"  {k} of {len(chunks)} requests, ${total:.2f} so far")
    print(f"AI read done: ${total:.2f}; {missing} left unlabeled (they fall back to the word rules)")
    return total


SITE_CACHE = OUT / f"vertical-sites-{DATE}.jsonl"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0 Safari/537.36"}


def _clean(s):
    s = re.sub(r"<[^>]+>", " ", s or "")
    s = re.sub(r"&amp;", "&", s)
    s = re.sub(r"&#0?39;|&rsquo;|&apos;", "'", s)
    return re.sub(r"\s+", " ", re.sub(r"&[a-z#0-9]+;", " ", s)).strip()[:300]


def fetch_site(dom):
    """The homepage's title, meta description, and first heading: one GET, https then http."""
    for scheme in ("https", "http"):
        try:
            r = requests.get(f"{scheme}://{dom}/", timeout=12, headers=UA, allow_redirects=True)
            html = r.text[:300000]
            break
        except requests.RequestException:
            html = None
    if not html:
        return {"domain": dom, "text": ""}
    title = re.search(r"<title[^>]*>(.*?)</title>", html, re.S | re.I)
    desc = (re.search(r"<meta[^>]+(?:name|property)=[\"'](?:og:)?description[\"'][^>]*content=[\"']([^\"']*)", html, re.I)
            or re.search(r"<meta[^>]+content=[\"']([^\"']*)[\"'][^>]*(?:name|property)=[\"'](?:og:)?description", html, re.I))
    h1 = re.search(r"<h1[^>]*>(.*?)</h1>", html, re.S | re.I)
    parts = [f"Site title: {_clean(title.group(1))}" if title else "", f"Description: {_clean(desc.group(1))}" if desc else "",
             f"Heading: {_clean(h1.group(1))}" if h1 else ""]
    return {"domain": dom, "text": ". ".join(p for p in parts if p and not p.endswith(": "))}


def site_cache():
    out = {}
    if SITE_CACHE.exists():
        for line in SITE_CACHE.read_text(encoding="utf-8").splitlines():
            rec = json.loads(line)
            out[rec["domain"]] = rec.get("text", "")
    return out


def crawl(domains):
    if not domains:
        return
    print(f"reading {len(domains)} homepages (title, description, first heading)")
    with ThreadPoolExecutor(max_workers=24) as ex, SITE_CACHE.open("a", encoding="utf-8") as fh:
        for k, rec in enumerate(ex.map(fetch_site, domains), 1):
            fh.write(json.dumps(rec) + "\n")
            if k % 250 == 0:
                print(f"  {k} of {len(domains)}")


def main():
    rows = pull("--cached" in sys.argv)
    reads, gtypes = site_reads(), google_types()
    todo = {}
    for r in rows:
        dom = norm_domain(str(r.get("Website") or ""))
        s = reads.get(dom)
        if s and s["verdict"] == "no" and not (r.get("Queue") and r["Queue"] <= CHECKED_LAST_QUEUE):
            todo[dom] = (str(r.get("Name") or "").strip(), s["focus"], s["reason"])
    done = ai_cache()
    pending = {d: v for d, v in todo.items() if d not in done}
    if pending:
        print(f"{len(pending)} 'not an integrator' site reads still need the AI read" +
              ("" if "--ai" in sys.argv else " (run with --ai; until then they use the word rules)"))
        if "--ai" in sys.argv:
            if "--limit" in sys.argv:  # a smoke test: the first N only
                n = int(sys.argv[sys.argv.index("--limit") + 1])
                pending = dict(sorted(pending.items())[:n])
            run_ai(pending)
            done = ai_cache()

    crawled = site_cache()

    def decide(r):
        dom = norm_domain(str(r.get("Website") or ""))
        s = reads.get(dom)
        if r.get("Queue") and r["Queue"] <= CHECKED_LAST_QUEUE:
            return "AV integrator", "checked by hand"
        if s and s["verdict"] == "yes":
            return "AV integrator", "site read: integrator"
        if s and s["verdict"] == "no" and done.get(dom) in LABELS:
            return done[dom], "site read + AI"
        label = by_words(r.get("Name"), r.get("Website"), s["title"] if s else "")
        if label:
            return label, "name words"
        t = gtypes.get(("d", dom)) or gtypes.get(("t", digits(r.get("Phone")))) or ""
        label = GOOGLE.get(t) or ("Not a contractor" if t in NOT_CONTRACTOR_TYPES else None)
        if label:
            return label, f"Google: {t}"
        if crawled.get(dom) and done.get(dom) in LABELS:
            return done[dom], "homepage + AI"
        return ("AV integrator" if r.get("Queue") else "Electrician"), "list default"

    # The rows that would fall to the list default: read their homepage's title and description, then the same AI read
    if "--crawl" in sys.argv:
        targets = sorted({norm_domain(str(r.get("Website") or "")) for r in rows if decide(r)[1] == "list default"} - {""})
        fetch = [d for d in targets if d not in crawled]
        print(f"list-default rows: {len(targets)} websites; {len(fetch)} not fetched yet")
        crawl(fetch)
        crawled = site_cache()
        todo2 = {d: (next((str(r.get("Name") or "") for r in rows if norm_domain(str(r.get("Website") or "")) == d), ""),
                     "", crawled[d]) for d in targets if crawled.get(d) and d not in done}
        if todo2 and "--ai" in sys.argv:
            run_ai(todo2)
            done = ai_cache()

    out, why = [], collections.Counter()
    for r in rows:
        label, rule = decide(r)
        why[rule.split(":")[0]] += 1
        out.append({"Company ID": r.get("Company ID"), "Name": r.get("Name"), "Website": r.get("Website"),
                    "Old": r.get("Vertical") or "", "New": label, "Rule": rule, "id": r["id"],
                    "List": "integrators" if r.get("Queue") else "electrician list",
                    "Status (typed)": r.get("Status (typed)") or ""})

    with (OUT / f"vertical-labels-{DATE}.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(out[0]))
        w.writeheader()
        w.writerows(out)
    print(f"\nrows: {len(out)}; decided by: {dict(why)}")
    for lst in ("integrators", "electrician list"):
        c = collections.Counter(o["New"] for o in out if o["List"] == lst)
        print(f"{lst}: " + ", ".join(f"{k} {v}" for k, v in c.most_common()))
    total = collections.Counter(o["New"] for o in out)
    print("all: " + ", ".join(f"{k} {v}" for k, v in total.most_common()))
    changed = [o for o in out if o["New"] != o["Old"]]
    print(f"labels that change: {len(changed)}; review file: projects/outreach/vertical-labels-{DATE}.csv")
    wrong = [o for o in out if re.search(r"wrong vert", o["Status (typed)"], re.I)]
    print(f"rows he marked 'wrong vertical' while calling: {len(wrong)}; now labeled: "
          + ", ".join(f"{k} {v}" for k, v in collections.Counter(o['New'] for o in wrong).most_common()))

    if "--write" not in sys.argv:
        print("dry run: nothing written to Airtable")
        return
    if "--cached" in sys.argv:
        sys.exit("--write pulls the table fresh; run it without --cached")
    before = OUT / f"vertical-before-{DATE}.json"
    if not before.exists():
        before.write_text(json.dumps({r["id"]: r.get("Vertical") for r in rows}), encoding="utf-8")
    s, made = airtable(), 0
    for i in range(0, len(changed), 10):
        chunk = [{"id": o["id"], "fields": {"Vertical": o["New"]}} for o in changed[i:i + 10]]
        r = call(s, "PATCH", f"{API}/{BASE}/{CONTACTS}", json={"records": chunk, "typecast": True})
        if r.status_code >= 300:
            print(f"stopped after {made}: {r.status_code} {r.text[:300]}")
            break
        made += len(chunk)
        if made % 1000 < 10:
            print(f"  {made} written")
    print(f"written {made} of {len(changed)}; old labels kept in {before.name}")


if __name__ == "__main__":
    main()
