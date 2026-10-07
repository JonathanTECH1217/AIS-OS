"""Organize Portal proposal CSV exports into one workbook and print a summary for the findings doc.

Input: Portal.io per-proposal CSV exports (one file per proposal; first data row is the proposal
summary, the rest are line items). Files are moved from the repo root into
projects/proposals-pha/raw/ on first run.

Output: projects/proposals-pha/pha-proposals-<DATE>.xlsx with sheets Proposals, Line items,
Categories, Brands, Areas, Labor rates, Suppliers, Revisions, Method, plus a JSON summary on stdout.

Usage: python scripts/proposals_analyze.py
"""
import csv, json, re, shutil, statistics, sys
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "projects" / "proposals-pha"
RAW = OUT_DIR / "raw"
DATE = "2026-09-09"
OUT_XLSX = OUT_DIR / f"pha-proposals-{DATE}.xlsx"

# ---------- classification ----------
# order matters: first match wins for an item
ITEM_RULES = [
    ("Shades and screens", r"shade|palladiom|roller screen|insect screen|magnatrack|villamater|progressive screens|vinyl panel|somfy|hunter douglas|drapery|motorized screen|retractable screen"),
    ("Home theater", r"projector|projection|\bscreen\b|surround|theat(?:er|re)|anthem|kaleidescape|severtson|stewart|fortress|\blcr\b|avm ?70|mca ?525"),
    ("Lighting control", r"lutron(?!.*shade)|ketra|homeworks|radiora|\bdmf\b|lumaris|proluxe|tape ?light|cove|dimmer|keypad|light(?:ing)? control|palladiom keypad|canvas light|linear track|fixture|landscape light|exterior light|colorbeam|\bwac\b|usai"),
    ("Control and automation", r"control ?4|c4-|\bc4\b|savant|josh\.?ai|josh ai|crestron|\belan\b|\burc\b|\brti\b|touch ?screen|touchpanel|touch panel|core ?[135]|ea-[135]|remote control|halo remote|neeo|programming"),
    ("Outdoor audio and video", r"mariner|coastal source|outdoor|exterior speaker|sunbrite|pool tv|lanai|patio|landscape speaker|garden"),
    ("Whole-home audio", r"sonance|sonos|speaker|\bamp\b|amplifier|subwoofer|\bsub\b|audio|james loudspeaker|origin acoustics|triad|episode|russound|autonomic|matrix|denon|marantz|yamaha|klipsch|\bkef\b|focal|bowers|paradigm|turntable"),
    ("Cameras", r"camera|cctv|\bnvr\b|\bdvr\b|hikvision|icrealtime|ic realtime|luma|axis|surveillance|protect|reolink|\bring\b"),
    ("Access control and gates", r"gate|access control|card reader|maglock|mag lock|electric strike|butterfly|intercom|push to exit|reader|door station|\bdsx\b|liftmaster|door lock|doorlock|deadbolt|strike|2n |akuvox|comelit|entry system|barrier arm"),
    ("Security and alarm", r"alarm|qolsys|\bdsc\b|resideo|honeywell|sensor|siren|panel|keypad|water valve|leak|smoke|doorbell|doorbird|monitoring"),
    ("Network and connectivity", r"ubiquiti|unifi|araknis|network|router|switch|access point|wi-?fi|starlink|star link|nextivity|cel-?fi|cellular|antenna|booster|fiber|cloud key|\bpoe\b|internet|modem|firewall|meraki|ruckus|pakedge|luxul"),
    ("TV and video", r"\btvs?\b|hdtv|television|samsung|\blg\b|sony|display|video wall|menu board|mount|mantel|mantle|sanus|chief|frame|hdmi|apple tv|roku|cable box|video service|directv|seura|séura|projector lift|tv lift"),
    ("Electrical", r"electric|panel upgrade|chandelier lift|outlet|breaker|generator|\bspan\b|enphase"),
    ("Wiring and infrastructure", r"cat ?6|cat ?5|coax|wire|cable|conduit|rack|enclosure|power conditioner|furman|wattbox|panamax|surge|\bups\b|red atom|middle atlantic|structured|pre-?wire|rough-?in|jack|plate|bracket|misc\. parts|installation"),
]
ITEM_RE = [(c, re.compile(p, re.I)) for c, p in ITEM_RULES]

COMMERCIAL_RE = re.compile(
    r"westin|ledo|\balta\b|\bnova\b|crossroads|carrollton|abners|aagaming|gaming|kiddie academy|library|kemira|marina|"
    r"boat yard|showroom|conference|training room|\boffice\b|parking lot|menu board|festival wine|the yards|juniper|"
    r"marlow|loading dock|amenity|pool gate|job site|pizza|seafood shack|hut|wine and spirits|meeting room|ballroom|"
    r"butterfly|\bdsx\b|card reader|maglock|push to exit|multifamily|apartment|restaurant|church|school|clinic|"
    r"warehouse|division 28|bulletin|kwanzaa|lobby|tenant|suite", re.I)
BUILDER_RE = re.compile(
    r"builders?\b|design build|for hci|bayview|gate one|\bgyc\b|hammer and nails|pinehurst|baldwin homes|back nine|"
    r"new construction|preliminary budget|budget proposal|residence for|project for|rough-?in|pre-?wire", re.I)
CHANGE_ORDER_RE = re.compile(r"change order|bulletin", re.I)
REV_RE = re.compile(r"\b(?:rev(?:ision)?\.?\s*#?\s*(\d+))", re.I)


def norm_title(t):
    t = t.lower()
    t = re.sub(r"\(\d+\)", " ", t)
    t = re.sub(r"copy of", " ", t)
    t = re.sub(r"\brev(?:ision)?\.?\s*#?\s*\d+\b", " ", t)
    t = re.sub(r"\brevised\b|\bupdated\b|\bproposal\b|\bnew\b|\bdate[d]?\b|\boption\b", " ", t)
    t = re.sub(r"\d[\d\-/]*\d", " ", t)  # dates and numbers
    t = re.sub(r"[^a-z& ]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def date_candidates(t):
    """Return possible dates found in a title. Ambiguous digit runs yield several candidates."""
    out = []
    for m in re.finditer(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})", t):
        mo, d, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y = y + 2000 if y < 100 else y
        try:
            out.append(date(y, mo, d))
        except ValueError:
            pass
    for m in re.finditer(r"(?<!\d)(\d{6,8})(?!\d)", t):
        s = m.group(1)
        combos = []
        if len(s) == 8:
            combos += [(s[0:2], s[2:4], s[4:8]), (s[0:2], s[2:4], s[4:6])]
        if len(s) == 7:
            combos += [(s[0:1], s[1:3], s[3:7]), (s[0:2], s[2:3], s[3:7])]
        if len(s) == 6:
            combos += [(s[0:2], s[2:4], s[4:6])]
        for mo, d, y in combos:
            try:
                mo, d, y = int(mo), int(d), int(y)
                y = y + 2000 if y < 100 else y
                if 2023 <= y <= 2026:
                    out.append(date(y, mo, d))
            except ValueError:
                pass
    return out


def fnum(v):
    try:
        return float(v) if v not in (None, "") else None
    except ValueError:
        return None


def load_file(path):
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        return None
    summary = None
    items = []
    for r in rows:
        if not (r.get("ItemType") or "").strip() and fnum(r.get("GrandTotal")) is not None:
            summary = r
        elif (r.get("ItemType") or "").strip():
            items.append(r)
    return summary, items


def classify_item(r, title):
    text = " ".join([r.get("Brand") or "", r.get("Model or Labor/Fee Name") or "", r.get("ShortDescription") or "",
                     r.get("Area") or ""])
    for cat, rx in ITEM_RE:
        if rx.search(text):
            return cat
    return None


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    RAW.mkdir(exist_ok=True)
    moved = 0
    for f in sorted({p.resolve() for p in list(ROOT.glob("Proposal #*.CSV")) + list(ROOT.glob("Proposal #*.csv"))}):
        if f.exists():
            shutil.move(str(f), str(RAW / f.name))
            moved += 1
    files = sorted({p.resolve() for p in list(RAW.glob("*.CSV")) + list(RAW.glob("*.csv"))})
    print(f"moved {moved} files; {len(files)} files in raw", file=sys.stderr)

    # one record per proposal number: keep the biggest file for that number
    by_num = {}
    for f in files:
        m = re.match(r"Proposal #(\d+)\s*-\s*(.*)\.csv$", f.name, re.I)
        if not m:
            continue
        num, title = int(m.group(1)), m.group(2).strip()
        title = re.sub(r"\s*\(\d+\)\s*$", "", title)
        if num not in by_num or f.stat().st_size > by_num[num]["size"]:
            by_num[num] = {"num": num, "title": title, "file": f, "size": f.stat().st_size}

    proposals, all_items = [], []
    empties = 0
    for num in sorted(by_num):
        rec = by_num[num]
        loaded = load_file(rec["file"])
        if not loaded or (loaded[0] is None and not loaded[1]):
            empties += 1
            continue
        summary, items = loaded
        title = rec["title"]
        cats = Counter()
        area_cat = {}
        parts_sell = labor_sell = fee_sell = 0.0
        parts_cost = labor_cost = 0.0
        parts_cost_known_sell = labor_cost_known_sell = 0.0
        brands = Counter()
        areas = Counter()
        labor_hours = Counter()
        n_items = 0
        for r in items:
            n_items += 1
            it = (r.get("ItemType") or "").strip()
            sell = fnum(r.get("SellTotal")) or 0.0
            cost = fnum(r.get("CostTotal"))
            area = (r.get("Area") or "").strip()
            brand = (r.get("Brand") or "").strip()
            if it == "Labor":
                labor_sell += sell
                name = (r.get("Model or Labor/Fee Name") or "").strip()
                if cost is not None:
                    labor_cost += cost
                    labor_cost_known_sell += sell
                labor_hours[(name, fnum(r.get("SellPrice")))] += fnum(r.get("AreaQty")) or 0
                cat = None
            elif it == "Fee":
                fee_sell += sell
                cat = None
            else:
                parts_sell += sell
                if cost is not None:
                    parts_cost += cost
                    parts_cost_known_sell += sell
                if brand and brand.lower() not in ("custom", "labor", "fee"):
                    brands[brand] += sell
                cat = classify_item(r, title)
                if cat:
                    cats[cat] += sell
                    area_cat.setdefault(area, Counter())[cat] += sell
            areas[area] += sell
            all_items.append({
                "Proposal": num, "Title": title, "Change order": r.get("Change Order") or "", "Area": area,
                "ItemType": it, "Brand": brand, "Model": r.get("Model or Labor/Fee Name") or "",
                "Description": r.get("ShortDescription") or "", "Qty": fnum(r.get("AreaQty")),
                "Sell price": fnum(r.get("SellPrice")), "Sell total": sell, "Cost total": cost,
                "Supplier": r.get("Supplier") or "", "Profit total": fnum(r.get("ProfitTotal")),
                "Margin %": fnum(r.get("ProfitPercent")), "Category": cat or ("Labor" if it == "Labor" else ""),
            })
        # labor follows the dominant category of its area; unclassified areas follow the proposal's top category
        for r in items:
            if (r.get("ItemType") or "").strip() == "Labor":
                area = (r.get("Area") or "").strip()
                sell = fnum(r.get("SellTotal")) or 0.0
                ac = area_cat.get(area)
                if ac:
                    cats[ac.most_common(1)[0][0]] += sell
                elif cats:
                    cats[cats.most_common(1)[0][0]] += sell
        grand = fnum(summary.get("GrandTotal")) if summary else None
        if grand is None:
            grand = parts_sell + labor_sell + fee_sell
        tax = fnum(summary.get("TotalTax")) if summary else None
        discount = fnum(summary.get("TotalPartsDiscount")) if summary else None
        primary = cats.most_common(1)[0][0] if cats else ""
        text = title + " " + " ".join(areas) + " " + " ".join(brands)
        commercial = bool(COMMERCIAL_RE.search(text)) and not re.search(r"residence|home|house", title, re.I)
        if re.search(r"westin|ledo|alta|nova|abners|gaming|kiddie|library|kemira|marina|boat yard|showroom|festival|the yards|juniper|marlow", title, re.I):
            commercial = True
        rev = REV_RE.search(title)
        proposals.append({
            "Proposal": num, "Title": title, "Norm title": norm_title(title),
            "Rev": int(rev.group(1)) if rev else 0, "Date candidates": date_candidates(title),
            "Grand total": round(grand, 2), "Parts": round(parts_sell, 2), "Labor": round(labor_sell, 2),
            "Fees": round(fee_sell, 2), "Tax": tax, "Discount": discount,
            "Parts cost": round(parts_cost, 2), "Parts sell w/ cost": round(parts_cost_known_sell, 2),
            "Labor cost": round(labor_cost, 2), "Labor sell w/ cost": round(labor_cost_known_sell, 2),
            "Items": n_items, "Areas": "; ".join(a for a in areas if a), "Area count": len([a for a in areas if a]),
            "Primary category": primary,
            "Category $": {k: round(v, 2) for k, v in cats.most_common()},
            "Top brands": "; ".join(f"{b} ${v:,.0f}" for b, v in brands.most_common(5)),
            "Brand $": dict(brands),
            "Segment": "Commercial" if commercial else "Residential",
            "Builder channel": bool(BUILDER_RE.search(title)) and not commercial,
            "Change order": bool(CHANGE_ORDER_RE.search(title)),
            "Labor hours": dict(labor_hours),
            "File": rec["file"].name,
        })
    print(f"{len(proposals)} proposals parsed, {empties} empty exports skipped", file=sys.stderr)

    # ---- dates: anchor on unambiguous single-candidate titles, interpolate the rest by proposal number
    anchors = sorted((p["Proposal"], p["Date candidates"][0]) for p in proposals if len(set(p["Date candidates"])) == 1)
    def interp(num):
        if not anchors:
            return None
        lo = max((a for a in anchors if a[0] <= num), default=None)
        hi = min((a for a in anchors if a[0] >= num), default=None)
        if lo and hi and hi[0] != lo[0]:
            frac = (num - lo[0]) / (hi[0] - lo[0])
            return lo[1] + timedelta(days=(hi[1] - lo[1]).days * frac)
        return (lo or hi)[1]
    for p in proposals:
        est = interp(p["Proposal"])
        cands = p["Date candidates"]
        if cands and est:
            p["Date"] = min(cands, key=lambda d: abs((d - est).days))
            p["Date source"] = "title"
        elif cands:
            p["Date"] = cands[0]; p["Date source"] = "title"
        else:
            p["Date"] = est; p["Date source"] = "estimated from proposal number"
        p["Month"] = p["Date"].strftime("%Y-%m") if p["Date"] else ""

    # ---- revisions: group by normalized title; latest number is the live version
    groups = defaultdict(list)
    for p in proposals:
        groups[p["Norm title"] or f"#{p['Proposal']}"].append(p)
    for key, ps in groups.items():
        ps.sort(key=lambda x: x["Proposal"])
        for i, p in enumerate(ps):
            p["Job key"] = key
            p["Revisions in group"] = len(ps)
            p["Status"] = "latest" if i == len(ps) - 1 else "superseded"
    latest = [p for p in proposals if p["Status"] == "latest" and p["Grand total"] > 0]

    # ---- rollups on latest versions
    def money_stats(vals):
        vals = sorted(vals)
        return {"count": len(vals), "total": round(sum(vals)), "median": round(statistics.median(vals)) if vals else 0,
                "mean": round(statistics.mean(vals)) if vals else 0, "max": round(max(vals)) if vals else 0}
    res = [p for p in latest if p["Segment"] == "Residential"]
    com = [p for p in latest if p["Segment"] == "Commercial"]
    builder = [p for p in res if p["Builder channel"]]
    non_builder = [p for p in res if not p["Builder channel"] and not p["Change order"]]
    buckets = Counter()
    bucket_dollars = Counter()
    for p in latest:
        g = p["Grand total"]
        b = "<$5k" if g < 5000 else "$5k-25k" if g < 25000 else "$25k-100k" if g < 100000 else "$100k+"
        buckets[b] += 1
        bucket_dollars[b] += g
    cat_rows = defaultdict(lambda: {"proposals": 0, "dollars": 0.0, "primary": 0, "primary_dollars": 0.0, "tickets": []})
    for p in latest:
        for c, v in p["Category $"].items():
            cat_rows[c]["proposals"] += 1
            cat_rows[c]["dollars"] += v
        if p["Primary category"]:
            cr = cat_rows[p["Primary category"]]
            cr["primary"] += 1
            cr["primary_dollars"] += p["Grand total"]
            cr["tickets"].append(p["Grand total"])
    brand_rows = defaultdict(lambda: {"proposals": 0, "dollars": 0.0, "big": 0, "sell_w_cost": 0.0, "cost": 0.0})
    for p in latest:
        for b, v in p["Brand $"].items():
            brand_rows[b]["proposals"] += 1
            brand_rows[b]["dollars"] += v
            if p["Grand total"] >= 25000:
                brand_rows[b]["big"] += 1
    for it in all_items:
        if it["ItemType"] == "Part" and it["Brand"] and it["Cost total"] is not None and it["Sell total"]:
            br = brand_rows[it["Brand"]]
            br["sell_w_cost"] += it["Sell total"]
            br["cost"] += it["Cost total"]
    area_counter = Counter()
    for p in res:
        for a in p["Areas"].split("; "):
            if a:
                area_counter[re.sub(r"\s+", " ", a.strip().lower())] += 1
    labor = Counter()
    labor_sell = Counter()
    for p in proposals:
        for (name, rate), hrs in p["Labor hours"].items():
            labor[(name.strip(), rate)] += hrs
    supplier = Counter()
    for it in all_items:
        if it["Supplier"]:
            supplier[it["Supplier"]] += it["Sell total"] or 0
    parts_sell_wc = sum(p["Parts sell w/ cost"] for p in latest)
    parts_cost = sum(p["Parts cost"] for p in latest)
    labor_sell_wc = sum(p["Labor sell w/ cost"] for p in latest)
    labor_cost = sum(p["Labor cost"] for p in latest)
    months = Counter()
    month_dollars = Counter()
    for p in latest:
        if p["Month"]:
            months[p["Month"]] += 1
            month_dollars[p["Month"]] += p["Grand total"]
    top = sorted(latest, key=lambda p: -p["Grand total"])[:25]
    multi_rev = sorted([ps[-1] for ps in groups.values() if len(ps) >= 3], key=lambda p: -p["Revisions in group"])[:15]

    # multi-system projects vs single-system, on residential latest versions
    def sys_count(p):
        return sum(1 for c, v in p["Category $"].items() if v >= 2000 and c != "Wiring and infrastructure")
    multi = [p for p in res if sys_count(p) >= 3]
    single = [p for p in res if sys_count(p) <= 1]
    combos = Counter()
    for p in multi:
        cats3 = tuple(sorted([c for c, v in p["Category $"].items() if v >= 2000 and c != "Wiring and infrastructure"])[:6])
        combos[cats3] += 1
    rider = Counter()  # categories that appear but are rarely primary
    for p in res:
        for c in p["Category $"]:
            if c != p["Primary category"] and p["Category $"][c] >= 1000:
                rider[c] += 1
    lutron_props = [p for p in res if "Lutron" in p["Brand $"] or "Ketra" in p["Brand $"]]
    big_res = [p for p in res if p["Grand total"] >= 25000]
    lutron_in_big = sum(1 for p in big_res if "Lutron" in p["Brand $"] or "Ketra" in p["Brand $"])
    lighting_props = [p for p in res if p["Category $"].get("Lighting control", 0) >= 5000]
    extra = {
        "res_multi_system": money_stats([p["Grand total"] for p in multi]),
        "res_single_system": money_stats([p["Grand total"] for p in single]),
        "top_combos": [(" + ".join(k), v) for k, v in combos.most_common(8)],
        "rider_categories": rider.most_common(8),
        "res_with_lutron_or_ketra": money_stats([p["Grand total"] for p in lutron_props]),
        "res_25k_plus": len(big_res), "res_25k_plus_with_lutron_or_ketra": lutron_in_big,
        "res_with_5k_plus_lighting": money_stats([p["Grand total"] for p in lighting_props]),
        "res_without_lighting": money_stats([p["Grand total"] for p in res if p["Category $"].get("Lighting control", 0) < 5000]),
        "electrical_passthrough": round(sum(v for p in latest for b, v in p["Brand $"].items() if "electrical project" in b.lower())),
    }

    summary = {
        "files": len(files), "unique_numbers": len(by_num), "parsed": len(proposals), "empty": empties,
        "latest_versions": len(latest), "superseded": len(proposals) - len(latest),
        "date_range": [min(p["Date"] for p in latest if p["Date"]).isoformat(), max(p["Date"] for p in latest if p["Date"]).isoformat()],
        "dates_from_title": sum(1 for p in latest if p["Date source"] == "title"),
        "all_latest": money_stats([p["Grand total"] for p in latest]),
        "residential": money_stats([p["Grand total"] for p in res]),
        "commercial": money_stats([p["Grand total"] for p in com]),
        "residential_builder": money_stats([p["Grand total"] for p in builder]),
        "residential_non_builder_non_co": money_stats([p["Grand total"] for p in non_builder]),
        "residential_change_orders": money_stats([p["Grand total"] for p in res if p["Change order"]]),
        "buckets": {b: {"count": buckets[b], "dollars": round(bucket_dollars[b])} for b in ["<$5k", "$5k-25k", "$25k-100k", "$100k+"]},
        "top10pct_share": round(sum(p["Grand total"] for p in sorted(latest, key=lambda p: -p["Grand total"])[: max(1, len(latest) // 10)]) / sum(p["Grand total"] for p in latest), 3),
        "categories": {c: {"in_proposals": v["proposals"], "dollars": round(v["dollars"]), "primary_count": v["primary"],
                           "primary_dollars": round(v["primary_dollars"]),
                           "primary_median_ticket": round(statistics.median(v["tickets"])) if v["tickets"] else 0}
                       for c, v in sorted(cat_rows.items(), key=lambda kv: -kv[1]["dollars"])},
        "brands_top": [{"brand": b, "proposals": v["proposals"], "dollars": round(v["dollars"]), "in_25k_plus": v["big"],
                        "margin_pct": round(100 * (v["sell_w_cost"] - v["cost"]) / v["sell_w_cost"], 1) if v["sell_w_cost"] else None}
                       for b, v in sorted(brand_rows.items(), key=lambda kv: -kv[1]["dollars"])[:30]],
        "areas_top": area_counter.most_common(30),
        "labor_rates": [{"name": n, "rate": r, "hours": round(h)} for (n, r), h in labor.most_common(12)],
        "parts_margin_pct": round(100 * (parts_sell_wc - parts_cost) / parts_sell_wc, 1) if parts_sell_wc else None,
        "labor_margin_pct": round(100 * (labor_sell_wc - labor_cost) / labor_sell_wc, 1) if labor_sell_wc else None,
        "parts_vs_labor_share": {"parts": round(sum(p["Parts"] for p in latest)), "labor": round(sum(p["Labor"] for p in latest)), "fees": round(sum(p["Fees"] for p in latest))},
        "suppliers_top": [(s, round(v)) for s, v in supplier.most_common(10)],
        "months": {m: {"count": months[m], "dollars": round(month_dollars[m])} for m in sorted(months)},
        "top25": [{"num": p["Proposal"], "title": p["Title"], "total": round(p["Grand total"]), "cat": p["Primary category"],
                   "segment": p["Segment"], "builder": p["Builder channel"], "date": p["Date"].isoformat() if p["Date"] else "",
                   "cats": p["Category $"], "brands": p["Top brands"]} for p in top],
        "multi_rev": [{"num": p["Proposal"], "title": p["Title"], "revs": p["Revisions in group"], "total": round(p["Grand total"])} for p in multi_rev],
        **extra,
    }

    # ---- workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Proposals"
    cols = ["Proposal", "Title", "Date", "Date source", "Status", "Revisions in group", "Grand total", "Parts", "Labor",
            "Fees", "Tax", "Discount", "Primary category", "Segment", "Builder channel", "Change order", "Items",
            "Area count", "Areas", "Top brands", "Category $", "File"]
    ws.append(cols)
    for p in sorted(proposals, key=lambda x: x["Proposal"]):
        row = []
        for c in cols:
            v = p.get(c)
            if c == "Date":
                v = v.isoformat() if v else ""
            elif c == "Category $":
                v = "; ".join(f"{k} ${val:,.0f}" for k, val in v.items())
            elif isinstance(v, bool):
                v = "yes" if v else ""
            row.append(v)
        ws.append(row)
    ws2 = wb.create_sheet("Line items")
    icols = ["Proposal", "Title", "Change order", "Area", "ItemType", "Category", "Brand", "Model", "Description", "Qty",
             "Sell price", "Sell total", "Cost total", "Profit total", "Margin %", "Supplier"]
    ws2.append(icols)
    for it in all_items:
        ws2.append([it.get(c) for c in icols])
    ws3 = wb.create_sheet("Categories")
    ws3.append(["Category", "Proposals containing it", "Dollars in latest proposals", "Proposals where primary",
                "Dollars of those proposals", "Median ticket when primary"])
    for c, v in summary["categories"].items():
        ws3.append([c, v["in_proposals"], v["dollars"], v["primary_count"], v["primary_dollars"], v["primary_median_ticket"]])
    ws4 = wb.create_sheet("Brands")
    ws4.append(["Brand", "Proposals", "Dollars (latest)", "In $25k+ proposals", "Margin % (where cost known)"])
    for b, v in sorted(brand_rows.items(), key=lambda kv: -kv[1]["dollars"]):
        ws4.append([b, v["proposals"], round(v["dollars"]), v["big"],
                    round(100 * (v["sell_w_cost"] - v["cost"]) / v["sell_w_cost"], 1) if v["sell_w_cost"] else None])
    ws5 = wb.create_sheet("Areas")
    ws5.append(["Area (residential proposals)", "Proposals"])
    for a, n in area_counter.most_common():
        ws5.append([a, n])
    ws6 = wb.create_sheet("Labor rates")
    ws6.append(["Labor line", "Sell rate", "Hours across all proposals"])
    for (n, r), h in labor.most_common():
        ws6.append([n, r, round(h, 1)])
    ws7 = wb.create_sheet("Suppliers")
    ws7.append(["Supplier", "Sell dollars"])
    for s, v in supplier.most_common():
        ws7.append([s, round(v)])
    ws8 = wb.create_sheet("Revisions")
    ws8.append(["Job key", "Revisions", "Proposal numbers", "Latest total"])
    for key, ps in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        if len(ps) > 1:
            ws8.append([key, len(ps), ", ".join(str(p["Proposal"]) for p in ps), ps[-1]["Grand total"]])
    ws9 = wb.create_sheet("Method")
    for line in [
        f"Built {DATE} from {len(files)} Portal CSV exports in projects/proposals-pha/raw/.",
        "One row per proposal number; when a number was exported twice the larger file was kept. Empty exports skipped.",
        "Grand total is Portal's summary row when present, else the sum of line items.",
        "Category per line item by brand and description keywords (see scripts/proposals_analyze.py); labor follows its area's parts.",
        "Primary category = category with the most dollars in the proposal.",
        "Segment: commercial when the title or areas name a business, multifamily, or hospitality site; residential otherwise.",
        "Builder channel: residential proposals whose title names a builder, a budget or preliminary proposal, or rough-in work.",
        "Revisions grouped by normalized title (rev numbers, dates, 'copy of' stripped). Latest number in a group = live version;",
        "rollups use latest versions only, so a job revised nine times counts once.",
        "Dates parsed from titles where present (Rev dates, MMDDYY runs); otherwise estimated by interpolating proposal number",
        "between dated neighbors. Treat estimated dates as month-level.",
        "No status column exists in the export: these are proposals sent, not jobs won.",
        "Sensitive: homeowner names and PHA's costs and margins. Keep local. Do not quote client names outside this folder.",
    ]:
        ws9.append([line])
    for w in wb.worksheets:
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
    wb.save(OUT_XLSX)
    summary["xlsx"] = str(OUT_XLSX)
    print(json.dumps(summary, indent=1, default=str))


if __name__ == "__main__":
    main()
