"""Turn a places_seed workbook into a call list CSV the CRM can dial from, one row per company.

The integrator list went through qualify_sites.py (the Opus read) before it became qualified-list-<DATE>.csv. A new
vertical can skip that pass and dial the raw sweep: this script writes the same columns the CRM's queue loader reads
(Name, City, State, Zip, County, Phone, Local time at noon ET, Rating, Review count, Reviews shown, Listings, Website,
Status, Source), collapses duplicate listings of one company (same website domain, else same phone) and keeps the
listing with the most reviews, drops rows merged in from the integrator lists (In Repository = yes), drops rows with
no phone, and orders by review count so the established shops come first.

Usage:
  python scripts/seed_to_list.py projects/outreach/places-seed-2026-09-24-electricians.xlsx electrician-list-2026-09-24.csv
  python scripts/seed_to_list.py <seed.xlsx> <out.csv> --keep-dmv        # keep DC, MD, VA rows (the integrator dial list drops them)

The CRM picks the list up by the glob in projects/crm/config.json ("verticals", queue_csv_glob) on its next start.
Jonathan, 2026-09-24: "Build me a list of electricians just like the list for integrators" and "separate this list
in the CRM from the integrators."
"""
import argparse
import csv
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from outreach_common import DMV_STATES, OUTREACH, norm_domain  # noqa: E402

COLS = ["Name", "City", "State", "Zip", "County", "Phone", "Local time at noon ET", "Rating", "Review count",
        "Reviews shown", "Listings", "Website", "Status", "Source"]


def digits(s):
    return "".join(ch for ch in str(s or "") if ch.isdigit())[-10:]


def key_for(website, phone):
    d = norm_domain(website)
    if d:
        return "d:" + d
    p = digits(phone)
    return "p:" + p if len(p) == 10 else ""


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("seed", help="places-seed-<DATE>.xlsx")
    ap.add_argument("out", help="output CSV name (written to projects/outreach unless a path is given)")
    ap.add_argument("--keep-dmv", action="store_true", help="keep DC, MD, VA rows")
    ap.add_argument("--sheet", default=None, help="sheet name (default: the first)")
    a = ap.parse_args(argv)

    wb = openpyxl.load_workbook(a.seed, read_only=True, data_only=True)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or "").strip() for h in next(rows)]
    ix = {h: i for i, h in enumerate(head)}

    def col(r, name):
        i = ix.get(name)
        return r[i] if i is not None and i < len(r) else None

    best = {}
    listings = {}
    dropped = {"in_repository": 0, "no_phone": 0, "dmv": 0, "no_key": 0}
    for r in rows:
        if not r or not col(r, "Name"):
            continue
        # Rows merged in from the integrator lists carry a Source list that is not the sweep's ("2026-09-05",
        # "monarc-os"); the Extras sheet marks them In Repository. Either way they are not this vertical.
        src = str(col(r, "Source list") or "")
        if str(col(r, "In Repository") or "").lower() == "yes" or (src and not src.startswith("places-")):
            dropped["in_repository"] += 1
            continue
        phone = col(r, "Phone") or ""
        if len(digits(phone)) != 10:
            dropped["no_phone"] += 1
            continue
        state = str(col(r, "State") or "").upper()
        if not a.keep_dmv and state in DMV_STATES:
            dropped["dmv"] += 1
            continue
        k = key_for(col(r, "Website"), phone)
        if not k:
            dropped["no_key"] += 1
            continue
        reviews = int(col(r, "Review count") or 0)
        listings[k] = listings.get(k, 0) + 1
        cur = best.get(k)
        if cur is None or reviews > cur["_reviews"]:
            best[k] = {
                "Name": str(col(r, "Name")).strip(), "City": col(r, "City") or "", "State": state, "Zip": col(r, "Zip") or "",
                "County": col(r, "County") or "", "Phone": phone, "Local time at noon ET": col(r, "Local Time @ 12pm ET") or "",
                "Rating": col(r, "Rating") if col(r, "Rating") is not None else "", "Review count": reviews,
                "Website": col(r, "Website") or "", "Status": "", "Source": "cold", "_reviews": reviews,
            }
    out_rows = []
    for k, r in best.items():
        r["Reviews shown"] = "yes" if r["_reviews"] > 0 else "no"
        r["Listings"] = listings[k]
        out_rows.append(r)
    out_rows.sort(key=lambda r: (-r["_reviews"], r["State"], r["Name"]))

    out = Path(a.out)
    if not out.is_absolute() and out.parent == Path("."):
        out = OUTREACH / out
    with open(out, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS, extrasaction="ignore")
        w.writeheader()
        for r in out_rows:
            w.writerow(r)
    print(f"{out}: {len(out_rows)} companies from {ws.max_row - 1} listings; dropped {dropped}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
